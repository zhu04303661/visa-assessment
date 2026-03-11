#!/usr/bin/env python3
"""
GTV签证文案制作系统 - API路由模块
提供材料收集、信息提取、框架构建等功能的REST API接口
"""

import os
import json
import sqlite3
from datetime import datetime
from typing import Dict, Any
from flask import Blueprint, request, jsonify, send_file
from werkzeug.utils import secure_filename
from pathlib import Path

# 加载环境变量（确保 MinIO 配置可用）
try:
    from dotenv import load_dotenv
    env_path = Path(__file__).parent.parent.parent / '.env.local'
    if env_path.exists():
        load_dotenv(env_path)
except ImportError:
    pass

# 创建 Blueprint
copywriting_bp = Blueprint('copywriting', __name__)

# 配置
UPLOAD_FOLDER = os.getenv("UPLOAD_FOLDER", "./uploads")
PROJECTS_PATH = os.getenv("COPYWRITING_PROJECTS_PATH", "./copywriting_projects")
CASES_PATH = os.getenv("CASE_LIBRARY_PATH", "./success_cases")
DB_PATH = os.getenv("COPYWRITING_DB_PATH", "./copywriting.db")

# 支持的文件类型
ALLOWED_EXTENSIONS = {
    'txt', 'pdf', 'doc', 'docx', 'md', 'json', 'rtf',
    'jpg', 'jpeg', 'png', 'gif', 'webp', 'bmp', 'tiff', 'heic',
    'csv', 'xlsx', 'xls',
    'zip'
}

# 全局服务实例（延迟初始化）
_services = {}

def _get_logger():
    """获取日志器"""
    if 'logger' not in _services:
        try:
            from utils.logger_config import setup_module_logger
            _services['logger'] = setup_module_logger("copywriting_routes", os.getenv("LOG_LEVEL", "INFO"))
        except ImportError:
            import logging
            logging.basicConfig(level=logging.INFO)
            _services['logger'] = logging.getLogger("copywriting_routes")
    return _services['logger']

def _init_services():
    """初始化所有服务（延迟加载）"""
    if _services.get('initialized'):
        return
    
    logger = _get_logger()
    
    try:
        from database.copywriting_database import CopywritingDatabase
        from processors.material_processor import MaterialProcessor
        from services.raw_material_manager import RawMaterialManager
        from processors.material_analyzer import MaterialAnalyzer
        from agents.content_extraction_agent import ContentExtractionAgent
        from agents.framework_building_agent import FrameworkBuildingAgent
        from services.copywriting_workflow import CopywritingWorkflow
        
        # 确保目录存在
        Path(UPLOAD_FOLDER).mkdir(parents=True, exist_ok=True)
        
        # 初始化服务
        _services['db'] = CopywritingDatabase(DB_PATH)
        logger.info(f"本地数据库初始化成功: {DB_PATH}")
        
        _services['material_processor'] = MaterialProcessor(UPLOAD_FOLDER)
        logger.info("材料处理器初始化成功")
        
        _services['raw_material_manager'] = RawMaterialManager(DB_PATH, UPLOAD_FOLDER)
        logger.info("原始材料管理器初始化成功")
        
        _services['material_analyzer'] = MaterialAnalyzer(DB_PATH)
        logger.info("材料分析器初始化成功")
        
        _services['content_extraction_agent'] = ContentExtractionAgent(DB_PATH, UPLOAD_FOLDER)
        logger.info("内容提取Agent初始化成功")
        
        _services['framework_building_agent'] = FrameworkBuildingAgent(DB_PATH)
        logger.info("框架构建Agent初始化成功")
        
        _services['workflow'] = CopywritingWorkflow(PROJECTS_PATH, CASES_PATH)
        _services['project_manager'] = _services['workflow'].project_manager
        _services['case_library'] = _services['workflow'].case_library
        _services['agent'] = _services['workflow'].agent
        logger.info("文案系统服务初始化成功")
        
        _services['initialized'] = True
        
    except Exception as e:
        logger.error(f"服务初始化失败: {e}")
        _services['initialized'] = False

def get_service(name):
    """获取服务实例"""
    _init_services()
    return _services.get(name)

def allowed_file(filename):
    """检查文件类型是否允许"""
    return '.' in filename and filename.rsplit('.', 1)[1].lower() in ALLOWED_EXTENSIONS


# ==================== 健康检查 ====================

@copywriting_bp.route('/health', methods=['GET'])
def health_check():
    """健康检查接口"""
    db = get_service('db')
    material_processor = get_service('material_processor')
    raw_material_manager = get_service('raw_material_manager')
    workflow = get_service('workflow')
    agent = get_service('agent')
    case_library = get_service('case_library')
    project_manager = get_service('project_manager')
    
    return jsonify({
        "status": "healthy",
        "timestamp": datetime.now().isoformat(),
        "service": "Copywriting API",
        "database": "SQLite (local)",
        "db_path": DB_PATH,
        "components": {
            "database": db is not None,
            "material_processor": material_processor is not None,
            "workflow": workflow is not None,
            "agent": agent is not None,
            "case_library": case_library is not None,
            "project_manager": project_manager is not None
        },
        "supported_file_types": list(ALLOWED_EXTENSIONS)
    })


# ==================== 项目管理 API ====================

@copywriting_bp.route('/projects', methods=['GET'])
def list_projects():
    """获取项目列表"""
    db = get_service('db')
    if not db:
        return jsonify({"success": False, "error": "服务未初始化"}), 500
    
    result = db.list_projects()
    return jsonify(result)


@copywriting_bp.route('/projects', methods=['POST'])
def create_project():
    """创建新项目"""
    db = get_service('db')
    if not db:
        return jsonify({"success": False, "error": "服务未初始化"}), 500
    
    data = request.get_json()
    client_name = data.get('client_name', 'Unknown')
    visa_type = data.get('visa_type', 'GTV')
    
    result = db.create_project(client_name, visa_type)
    return jsonify(result)


@copywriting_bp.route('/projects/<project_id>', methods=['GET'])
def get_project(project_id):
    """获取项目详情"""
    db = get_service('db')
    if not db:
        return jsonify({"success": False, "error": "服务未初始化"}), 500
    
    result = db.get_project(project_id=project_id)
    return jsonify(result)


@copywriting_bp.route('/projects/<project_id>', methods=['DELETE'])
def delete_project(project_id):
    """删除项目"""
    db = get_service('db')
    if not db:
        return jsonify({"success": False, "error": "服务未初始化"}), 500
    
    result = db.delete_project(project_id)
    return jsonify(result)


@copywriting_bp.route('/projects/<project_id>/workflow', methods=['GET'])
def get_workflow_status(project_id):
    """获取项目工作流状态（基于实际数据判断各步骤完成情况）"""
    import sqlite3
    logger = _get_logger()
    db_path = _get_db_path()
    
    try:
        conn = sqlite3.connect(db_path)
        conn.row_factory = sqlite3.Row
        cursor = conn.cursor()
        
        stages = {}
        
        # 1. 材料收集 - 检查 material_files 表
        cursor.execute('''
            SELECT COUNT(*) as count FROM material_files WHERE project_id = ?
        ''', (project_id,))
        material_count = cursor.fetchone()['count']
        
        if material_count > 0:
            stages['1_collect'] = {
                'status': 'completed',
                'message': f'已收集 {material_count} 个材料文件',
                'count': material_count
            }
        else:
            stages['1_collect'] = {
                'status': 'pending',
                'message': '待收集材料',
                'count': 0
            }
        
        # 2. 材料分析 - 检查 extracted_contents 和 content_classifications
        cursor.execute('''
            SELECT COUNT(*) as count FROM extracted_contents WHERE project_id = ?
        ''', (project_id,))
        extracted_count = cursor.fetchone()['count']
        
        cursor.execute('''
            SELECT COUNT(*) as count FROM content_classifications WHERE project_id = ?
        ''', (project_id,))
        classified_count = cursor.fetchone()['count']
        
        if classified_count > 0:
            stages['2_analyze'] = {
                'status': 'completed',
                'message': f'已提取 {extracted_count} 条内容，已分类 {classified_count} 条证据',
                'extracted': extracted_count,
                'classified': classified_count
            }
        elif extracted_count > 0:
            stages['2_analyze'] = {
                'status': 'in_progress',
                'message': f'已提取 {extracted_count} 条内容，待分类',
                'extracted': extracted_count,
                'classified': 0
            }
        else:
            stages['2_analyze'] = {
                'status': 'pending',
                'message': '待分析材料',
                'extracted': 0,
                'classified': 0
            }
        
        # 3. GTV框架 - 检查 gtv_frameworks 表
        cursor.execute('''
            SELECT COUNT(*) as count FROM gtv_frameworks WHERE project_id = ?
        ''', (project_id,))
        framework_count = cursor.fetchone()['count']
        
        if framework_count > 0:
            stages['3_framework'] = {
                'status': 'completed',
                'message': 'GTV框架已构建',
                'count': framework_count
            }
        else:
            stages['3_framework'] = {
                'status': 'pending',
                'message': '待构建GTV框架',
                'count': 0
            }
        
        # 4. 案例匹配 - 检查 case_matches 表
        cursor.execute('''
            SELECT COUNT(*) as count FROM case_matches WHERE project_id = ?
        ''', (project_id,))
        match_count = cursor.fetchone()['count']
        
        if match_count > 0:
            stages['4_match'] = {
                'status': 'completed',
                'message': f'已匹配 {match_count} 个案例',
                'count': match_count
            }
        else:
            stages['4_match'] = {
                'status': 'pending',
                'message': '待匹配案例',
                'count': 0
            }
        
        # 5. 文案生成 - 检查 generated_documents 表
        cursor.execute('''
            SELECT COUNT(*) as count FROM generated_documents WHERE project_id = ?
        ''', (project_id,))
        doc_count = cursor.fetchone()['count']
        
        if doc_count > 0:
            stages['5_generate'] = {
                'status': 'completed',
                'message': f'已生成 {doc_count} 份文档',
                'count': doc_count
            }
        else:
            stages['5_generate'] = {
                'status': 'pending',
                'message': '待生成文案',
                'count': 0
            }
        
        # 6. 内容优化 & 7. 最终审核 - 暂时都设为 pending
        stages['6_optimize'] = {
            'status': 'pending',
            'message': '待优化内容'
        }
        stages['7_review'] = {
            'status': 'pending',
            'message': '待最终审核'
        }
        
        conn.close()
        
        return jsonify({
            "success": True,
            "data": {
                "project_id": project_id,
                "stages": stages
            }
        })
        
    except Exception as e:
        logger.error(f"获取工作流状态失败: {e}")
        import traceback
        traceback.print_exc()
        return jsonify({"success": False, "error": str(e)}), 500


@copywriting_bp.route('/projects/<project_id>/materials', methods=['GET'])
def get_project_materials(project_id):
    """获取项目材料"""
    db = get_service('db')
    project_manager = get_service('project_manager')
    
    if db:
        result = db.get_raw_materials(project_id)
        return jsonify(result)
    
    if not project_manager:
        return jsonify({"success": False, "error": "服务未初始化"}), 500
    
    result = project_manager.get_raw_materials(project_id)
    return jsonify(result)


# ==================== 材料收集 API ====================

@copywriting_bp.route('/material-collection/debug', methods=['GET'])
def debug_minio_status():
    """调试：检查 MinIO 状态"""
    import os
    from services import raw_material_manager as rmm_module
    
    # 强制重新初始化（用于调试）
    force_reinit = request.args.get('reinit', 'false').lower() == 'true'
    if force_reinit:
        _services['initialized'] = False
        if 'raw_material_manager' in _services:
            del _services['raw_material_manager']
    
    raw_material_manager = get_service('raw_material_manager')
    
    if not raw_material_manager:
        return jsonify({
            "success": False, 
            "error": "服务未初始化",
            "MINIO_IMPORT_OK": getattr(rmm_module, 'MINIO_IMPORT_OK', 'undefined'),
            "env": {
                "MINIO_ENDPOINT": os.getenv("MINIO_ENDPOINT"),
                "MINIO_ACCESS_KEY": os.getenv("MINIO_ACCESS_KEY"),
            }
        }), 500
    
    minio_manager = raw_material_manager.minio_manager
    return jsonify({
        "success": True,
        "MINIO_IMPORT_OK": getattr(rmm_module, 'MINIO_IMPORT_OK', 'undefined'),
        "use_minio": raw_material_manager.use_minio,
        "minio_enabled": minio_manager.is_enabled() if minio_manager else False,
        "minio_manager_exists": minio_manager is not None,
        "minio_endpoint": minio_manager.endpoint if minio_manager else None,
        "env": {
            "MINIO_ENDPOINT": os.getenv("MINIO_ENDPOINT"),
            "MINIO_ACCESS_KEY": os.getenv("MINIO_ACCESS_KEY"),
            "MINIO_SECRET_KEY": os.getenv("MINIO_SECRET_KEY", "")[:5] + "..." if os.getenv("MINIO_SECRET_KEY") else None,
        }
    })


@copywriting_bp.route('/material-collection/categories', methods=['GET'])
def get_material_categories():
    """获取材料分类结构"""
    raw_material_manager = get_service('raw_material_manager')
    if not raw_material_manager:
        return jsonify({"success": False, "error": "服务未初始化"}), 500
    
    result = raw_material_manager.get_material_categories()
    return jsonify(result)


@copywriting_bp.route('/material-collection/categories', methods=['PUT'])
def update_material_categories():
    """更新材料分类配置（保存到数据库）"""
    logger = _get_logger()
    
    data = request.get_json()
    categories = data.get('categories')
    
    if not categories:
        return jsonify({"success": False, "error": "缺少分类数据"}), 400
    
    try:
        # 保存到数据库
        from services.raw_material_manager import save_material_categories_to_db
        success = save_material_categories_to_db(categories)
        
        if success:
            logger.info(f"材料分类配置已保存到数据库")
            return jsonify({"success": True, "message": "分类配置已保存"})
        else:
            return jsonify({"success": False, "error": "保存失败"}), 500
    except Exception as e:
        logger.error(f"保存分类配置失败: {e}")
        return jsonify({"success": False, "error": str(e)}), 500


@copywriting_bp.route('/material-collection/forms', methods=['GET'])
def get_form_templates():
    """获取所有表单模板"""
    raw_material_manager = get_service('raw_material_manager')
    if not raw_material_manager:
        return jsonify({"success": False, "error": "服务未初始化"}), 500
    
    result = raw_material_manager.get_all_form_templates()
    return jsonify(result)


@copywriting_bp.route('/material-collection/forms/<form_type>', methods=['GET'])
def get_form_template(form_type):
    """获取特定表单模板"""
    raw_material_manager = get_service('raw_material_manager')
    if not raw_material_manager:
        return jsonify({"success": False, "error": "服务未初始化"}), 500
    
    result = raw_material_manager.get_form_template(form_type)
    return jsonify(result)


@copywriting_bp.route('/material-collection/templates/<form_type>/download', methods=['GET'])
def download_form_template(form_type):
    """下载采集表模板"""
    raw_material_manager = get_service('raw_material_manager')
    if not raw_material_manager:
        return jsonify({"success": False, "error": "服务未初始化"}), 500
    
    result = raw_material_manager.generate_form_template(form_type)
    
    if result.get("success"):
        file_path = result["data"]["file_path"]
        filename = result["data"]["filename"]
        return send_file(file_path, as_attachment=True, download_name=filename)
    else:
        return jsonify(result), 400


@copywriting_bp.route('/projects/<project_id>/material-collection', methods=['GET'])
def get_project_material_status(project_id):
    """获取项目材料收集状态"""
    raw_material_manager = get_service('raw_material_manager')
    if not raw_material_manager:
        return jsonify({"success": False, "error": "服务未初始化"}), 500
    
    result = raw_material_manager.get_collection_status(project_id)
    return jsonify(result)


@copywriting_bp.route('/projects/<project_id>/material-collection/init', methods=['POST'])
def init_project_materials(project_id):
    """初始化项目材料收集清单"""
    raw_material_manager = get_service('raw_material_manager')
    if not raw_material_manager:
        return jsonify({"success": False, "error": "服务未初始化"}), 500
    
    result = raw_material_manager.init_project_materials(project_id)
    return jsonify(result)


@copywriting_bp.route('/projects/<project_id>/material-collection/upload', methods=['POST'])
def upload_raw_material(project_id):
    """上传原始材料文件（支持 MinIO 存储）"""
    logger = _get_logger()
    raw_material_manager = get_service('raw_material_manager')
    
    if not raw_material_manager:
        return jsonify({"success": False, "error": "服务未初始化"}), 500
    
    if 'file' not in request.files:
        return jsonify({"success": False, "error": "未找到文件"}), 400
    
    file = request.files['file']
    category_id = request.form.get('category_id')
    item_id = request.form.get('item_id')
    description = request.form.get('description', '')
    
    if not category_id or not item_id:
        return jsonify({"success": False, "error": "缺少分类信息"}), 400
    
    if file.filename == '':
        return jsonify({"success": False, "error": "文件名为空"}), 400
    
    # 使用中文原始文件名
    original_filename = file.filename
    filename = secure_filename(file.filename)
    file_type = filename.rsplit('.', 1)[1].lower() if '.' in filename else ''
    
    # 读取文件数据
    file_data = file.read()
    file_size = len(file_data)
    
    # 使用 upload_material_bytes 直接上传到 MinIO（如果可用）
    result = raw_material_manager.upload_material_bytes(
        project_id=project_id,
        category_id=category_id,
        item_id=item_id,
        file_data=file_data,
        file_name=original_filename,
        file_size=file_size,
        file_type=file_type,
        description=description
    )
    
    logger.info(f"文件上传结果: {result.get('storage_type', 'unknown')}, minio_url={result.get('object_url', 'N/A')[:50] if result.get('object_url') else 'N/A'}")
    
    return jsonify(result)


@copywriting_bp.route('/projects/<project_id>/material-collection/files/<int:file_id>', methods=['DELETE'])
def delete_material_file(project_id, file_id):
    """删除材料文件"""
    raw_material_manager = get_service('raw_material_manager')
    if not raw_material_manager:
        return jsonify({"success": False, "error": "服务未初始化"}), 500
    
    result = raw_material_manager.delete_material_file(file_id)
    return jsonify(result)


@copywriting_bp.route('/projects/<project_id>/materials/<int:file_id>/tags', methods=['PUT'])
def update_material_tags(project_id, file_id):
    """更新材料文件的分类标签"""
    raw_material_manager = get_service('raw_material_manager')
    if not raw_material_manager:
        return jsonify({"success": False, "error": "服务未初始化"}), 500
    
    data = request.get_json()
    tags = data.get('tags', [])
    
    # 目前只支持单标签，取最后一个标签（最新添加的）
    if tags:
        category_id = tags[-1].get('category_id', '')
        item_id = tags[-1].get('item_id', '')
    else:
        category_id = ''
        item_id = ''
    
    result = raw_material_manager.update_material_tags(file_id, category_id, item_id)
    return jsonify(result)


# ==================== 内容提取 API ====================

@copywriting_bp.route('/projects/<project_id>/extraction/start', methods=['POST'])
def start_content_extraction(project_id):
    """启动内容提取任务"""
    content_extraction_agent = get_service('content_extraction_agent')
    if not content_extraction_agent:
        return jsonify({"success": False, "error": "内容提取Agent未初始化"}), 500
    
    result = content_extraction_agent.extract_and_classify(project_id)
    return jsonify(result)


@copywriting_bp.route('/projects/<project_id>/extraction/status', methods=['GET'])
def get_extraction_status(project_id):
    """获取提取进度"""
    content_extraction_agent = get_service('content_extraction_agent')
    if not content_extraction_agent:
        return jsonify({"success": False, "error": "内容提取Agent未初始化"}), 500
    
    result = content_extraction_agent.get_extraction_progress(project_id)
    return jsonify(result)


@copywriting_bp.route('/projects/<project_id>/extraction/results', methods=['GET'])
def get_extraction_results(project_id):
    """获取提取结果"""
    content_extraction_agent = get_service('content_extraction_agent')
    if not content_extraction_agent:
        return jsonify({"success": False, "error": "内容提取Agent未初始化"}), 500
    
    result = content_extraction_agent.get_classified_content(project_id)
    return jsonify(result)


# ==================== 框架构建 API ====================

@copywriting_bp.route('/projects/<project_id>/framework/build', methods=['POST'])
def build_framework(project_id):
    """构建申请框架"""
    framework_building_agent = get_service('framework_building_agent')
    if not framework_building_agent:
        return jsonify({"success": False, "error": "框架构建Agent未初始化"}), 500
    
    result = framework_building_agent.build_framework(project_id)
    return jsonify(result)


@copywriting_bp.route('/projects/<project_id>/framework', methods=['GET'])
def get_framework(project_id):
    """获取框架内容"""
    framework_building_agent = get_service('framework_building_agent')
    if not framework_building_agent:
        return jsonify({"success": False, "error": "框架构建Agent未初始化"}), 500
    
    result = framework_building_agent.get_framework(project_id)
    return jsonify(result)


@copywriting_bp.route('/projects/<project_id>/framework', methods=['PUT'])
def update_framework(project_id):
    """更新框架内容"""
    framework_building_agent = get_service('framework_building_agent')
    if not framework_building_agent:
        return jsonify({"success": False, "error": "框架构建Agent未初始化"}), 500
    
    data = request.get_json()
    result = framework_building_agent.update_framework(project_id, data)
    return jsonify(result)


# ==================== 文件预览 API ====================

@copywriting_bp.route('/files/preview/<int:file_id>', methods=['GET'])
def preview_file_by_id(file_id):
    """
    通过文件ID预览文件
    
    使用统一文件存储接口，自动支持任何存储后端
    """
    import sqlite3
    from flask import Response
    from urllib.parse import quote
    from database.file_storage import get_file_from_db_record, get_file_storage
    
    logger = _get_logger()
    db_path = DB_PATH
    
    try:
        conn = sqlite3.connect(db_path)
        conn.row_factory = sqlite3.Row
        cursor = conn.cursor()
        
        # 查询文件信息
        cursor.execute('''
            SELECT file_path, file_name, file_type, storage_type, object_bucket, object_key 
            FROM material_files WHERE id = ?
        ''', (file_id,))
        row = cursor.fetchone()
        conn.close()
        
        if not row:
            return jsonify({"success": False, "error": "文件不存在"}), 404
        
        file_name = row['file_name']
        
        # 使用统一存储接口获取文件内容
        record = dict(row)
        content = get_file_from_db_record(record)
        
        if content is None:
            return jsonify({"success": False, "error": "文件内容不可用"}), 404
        
        # 获取 MIME 类型
        storage = get_file_storage()
        content_type = storage.get_content_type(file_name)
        
        # 对文件名进行 URL 编码
        encoded_filename = quote(file_name, safe='')
        
        return Response(
            content,
            mimetype=content_type,
            headers={
                'Content-Disposition': f"inline; filename*=UTF-8''{encoded_filename}",
                'Access-Control-Allow-Origin': '*'
            }
        )
        
    except Exception as e:
        logger.error(f"文件预览失败: {e}")
        import traceback
        traceback.print_exc()
        return jsonify({"success": False, "error": str(e)}), 500


@copywriting_bp.route('/files/preview', methods=['GET'])
def preview_file():
    """预览文件内容"""
    logger = _get_logger()
    file_path = request.args.get('path')
    
    if not file_path:
        return jsonify({"success": False, "error": "缺少文件路径"}), 400
    
    # 安全检查
    if '..' in file_path:
        return jsonify({"success": False, "error": "无效的文件路径"}), 400
    
    # 检查文件是否存在
    if not os.path.exists(file_path):
        # 尝试在 uploads 目录下查找
        if not file_path.startswith('/'):
            alt_path = os.path.join(UPLOAD_FOLDER, file_path)
            if os.path.exists(alt_path):
                file_path = alt_path
            else:
                return jsonify({"success": False, "error": f"文件不存在: {file_path}"}), 404
        else:
            return jsonify({"success": False, "error": f"文件不存在: {file_path}"}), 404
    
    try:
        return send_file(file_path, as_attachment=False)
    except Exception as e:
        logger.error(f"文件预览失败: {e}")
        return jsonify({"success": False, "error": str(e)}), 500


@copywriting_bp.route('/files/download', methods=['GET'])
def download_file():
    """下载文件"""
    logger = _get_logger()
    file_path = request.args.get('path')
    
    if not file_path:
        return jsonify({"success": False, "error": "缺少文件路径"}), 400
    
    if '..' in file_path:
        return jsonify({"success": False, "error": "无效的文件路径"}), 400
    
    if not os.path.exists(file_path):
        if not file_path.startswith('/'):
            alt_path = os.path.join(UPLOAD_FOLDER, file_path)
            if os.path.exists(alt_path):
                file_path = alt_path
            else:
                return jsonify({"success": False, "error": f"文件不存在"}), 404
        else:
            return jsonify({"success": False, "error": f"文件不存在"}), 404
    
    try:
        return send_file(file_path, as_attachment=True)
    except Exception as e:
        logger.error(f"文件下载失败: {e}")
        return jsonify({"success": False, "error": str(e)}), 500


# ==================== 项目上下文 API ====================

@copywriting_bp.route('/projects/<project_id>/context', methods=['GET'])
def get_project_context(project_id):
    """获取项目完整上下文"""
    logger = _get_logger()
    try:
        content_extraction_agent = get_service('content_extraction_agent')
        if not content_extraction_agent:
            return jsonify({"success": False, "error": "内容提取器未初始化"}), 500
        
        with_sources = request.args.get('with_sources', 'false').lower() == 'true'
        result = content_extraction_agent.get_project_context(project_id, with_sources=with_sources)
        return jsonify(result)
        
    except Exception as e:
        logger.error(f"获取上下文失败: {e}")
        return jsonify({"success": False, "error": str(e)}), 500


@copywriting_bp.route('/projects/<project_id>/outline', methods=['GET'])
def get_content_outline(project_id):
    """获取项目的内容大纲"""
    logger = _get_logger()
    try:
        content_extraction_agent = get_service('content_extraction_agent')
        if not content_extraction_agent:
            return jsonify({"success": False, "error": "内容提取器未初始化"}), 500
        
        result = content_extraction_agent.get_content_outline(project_id)
        return jsonify(result)
        
    except Exception as e:
        logger.error(f"获取大纲失败: {e}")
        return jsonify({"success": False, "error": str(e)}), 500


@copywriting_bp.route('/projects/<project_id>/outline/generate', methods=['POST'])
def generate_content_outline(project_id):
    """生成内容大纲"""
    logger = _get_logger()
    try:
        content_extraction_agent = get_service('content_extraction_agent')
        if not content_extraction_agent:
            return jsonify({"success": False, "error": "内容提取器未初始化"}), 500
        
        result = content_extraction_agent.generate_outline(project_id)
        return jsonify(result)
        
    except Exception as e:
        logger.error(f"生成大纲失败: {e}")
        return jsonify({"success": False, "error": str(e)}), 500


@copywriting_bp.route('/projects/<project_id>/extract', methods=['POST'])
def extract_project_content(project_id):
    """执行内容提取"""
    logger = _get_logger()
    try:
        content_extraction_agent = get_service('content_extraction_agent')
        if not content_extraction_agent:
            return jsonify({"success": False, "error": "内容提取器未初始化"}), 500
        
        # 使用正确的方法名
        result = content_extraction_agent.extract_project_files(project_id)
        return jsonify(result)
        
    except Exception as e:
        logger.error(f"内容提取失败: {e}")
        return jsonify({"success": False, "error": str(e)}), 500


@copywriting_bp.route('/projects/<project_id>/extraction-logs', methods=['GET'])
def get_extraction_logs(project_id):
    """获取内容提取日志"""
    logger = _get_logger()
    try:
        content_extraction_agent = get_service('content_extraction_agent')
        if not content_extraction_agent:
            return jsonify({"success": False, "error": "内容提取器未初始化"}), 500
        
        limit = request.args.get('limit', 50, type=int)
        result = content_extraction_agent.get_extraction_logs(project_id, limit)
        return jsonify(result)
        
    except Exception as e:
        logger.error(f"获取提取日志失败: {e}")
        return jsonify({"success": False, "error": str(e)}), 500


@copywriting_bp.route('/projects/<project_id>/deduplicate', methods=['POST'])
def deduplicate_content(project_id):
    """内容去重"""
    logger = _get_logger()
    try:
        content_extraction_agent = get_service('content_extraction_agent')
        if not content_extraction_agent:
            return jsonify({"success": False, "error": "内容提取器未初始化"}), 500
        
        result = content_extraction_agent.deduplicate_content(project_id)
        return jsonify(result)
        
    except Exception as e:
        logger.error(f"内容去重失败: {e}")
        return jsonify({"success": False, "error": str(e)}), 500


@copywriting_bp.route('/projects/<project_id>/classify', methods=['POST'])
def classify_content(project_id):
    """内容分类"""
    logger = _get_logger()
    try:
        content_extraction_agent = get_service('content_extraction_agent')
        if not content_extraction_agent:
            return jsonify({"success": False, "error": "内容提取器未初始化"}), 500
        
        result = content_extraction_agent.classify_content(project_id)
        return jsonify(result)
        
    except Exception as e:
        logger.error(f"内容分类失败: {e}")
        return jsonify({"success": False, "error": str(e)}), 500


@copywriting_bp.route('/projects/<project_id>/classifications', methods=['GET'])
def get_classifications(project_id):
    """获取内容分类结果"""
    logger = _get_logger()
    try:
        content_extraction_agent = get_service('content_extraction_agent')
        if not content_extraction_agent:
            return jsonify({"success": False, "error": "内容提取器未初始化"}), 500
        
        result = content_extraction_agent.get_classifications(project_id)
        return jsonify(result)
        
    except Exception as e:
        logger.error(f"获取分类结果失败: {e}")
        return jsonify({"success": False, "error": str(e)}), 500


@copywriting_bp.route('/projects/<project_id>/classification-summary', methods=['GET'])
def get_classification_summary(project_id):
    """获取分类汇总"""
    logger = _get_logger()
    try:
        content_extraction_agent = get_service('content_extraction_agent')
        if not content_extraction_agent:
            return jsonify({"success": False, "error": "内容提取器未初始化"}), 500
        
        result = content_extraction_agent.get_classification_summary(project_id)
        return jsonify(result)
        
    except Exception as e:
        logger.error(f"获取分类汇总失败: {e}")
        return jsonify({"success": False, "error": str(e)}), 500


@copywriting_bp.route('/projects/<project_id>/classification-progress', methods=['GET'])
def get_classification_progress(project_id):
    """获取分类进度"""
    logger = _get_logger()
    try:
        content_extraction_agent = get_service('content_extraction_agent')
        if not content_extraction_agent:
            return jsonify({"success": False, "error": "内容提取器未初始化"}), 500
        
        result = content_extraction_agent.get_classification_progress(project_id)
        return jsonify(result)
        
    except Exception as e:
        logger.error(f"获取分类进度失败: {e}")
        return jsonify({"success": False, "error": str(e)}), 500


@copywriting_bp.route('/projects/<project_id>/classifications/<int:classification_id>', methods=['PUT'])
def update_classification(project_id, classification_id):
    """更新分类"""
    logger = _get_logger()
    try:
        content_extraction_agent = get_service('content_extraction_agent')
        if not content_extraction_agent:
            return jsonify({"success": False, "error": "内容提取器未初始化"}), 500
        
        data = request.get_json()
        result = content_extraction_agent.update_classification(project_id, classification_id, data)
        return jsonify(result)
        
    except Exception as e:
        logger.error(f"更新分类失败: {e}")
        return jsonify({"success": False, "error": str(e)}), 500


@copywriting_bp.route('/projects/<project_id>/classifications/<int:classification_id>', methods=['DELETE'])
def delete_classification(project_id, classification_id):
    """删除分类"""
    logger = _get_logger()
    try:
        content_extraction_agent = get_service('content_extraction_agent')
        if not content_extraction_agent:
            return jsonify({"success": False, "error": "内容提取器未初始化"}), 500
        
        result = content_extraction_agent.delete_classification(project_id, classification_id)
        return jsonify(result)
        
    except Exception as e:
        logger.error(f"删除分类失败: {e}")
        return jsonify({"success": False, "error": str(e)}), 500


@copywriting_bp.route('/projects/<project_id>/classifications', methods=['POST'])
def add_classification(project_id):
    """添加分类"""
    logger = _get_logger()
    try:
        content_extraction_agent = get_service('content_extraction_agent')
        if not content_extraction_agent:
            return jsonify({"success": False, "error": "内容提取器未初始化"}), 500
        
        data = request.get_json()
        result = content_extraction_agent.add_classification(project_id, data)
        return jsonify(result)
        
    except Exception as e:
        logger.error(f"添加分类失败: {e}")
        return jsonify({"success": False, "error": str(e)}), 500


@copywriting_bp.route('/projects/<project_id>/content-blocks', methods=['GET'])
def get_content_blocks(project_id):
    """获取内容块"""
    logger = _get_logger()
    try:
        content_extraction_agent = get_service('content_extraction_agent')
        if not content_extraction_agent:
            return jsonify({"success": False, "error": "内容提取器未初始化"}), 500
        
        result = content_extraction_agent.get_content_blocks(project_id)
        return jsonify(result)
        
    except Exception as e:
        logger.error(f"获取内容块失败: {e}")
        return jsonify({"success": False, "error": str(e)}), 500


@copywriting_bp.route('/projects/<project_id>/content/search', methods=['GET'])
def search_content(project_id):
    """搜索内容"""
    logger = _get_logger()
    try:
        content_extraction_agent = get_service('content_extraction_agent')
        if not content_extraction_agent:
            return jsonify({"success": False, "error": "内容提取器未初始化"}), 500
        
        query = request.args.get('q', '')
        category = request.args.get('category', '')
        result = content_extraction_agent.search_content(project_id, query, category)
        return jsonify(result)
        
    except Exception as e:
        logger.error(f"搜索内容失败: {e}")
        return jsonify({"success": False, "error": str(e)}), 500


# ==================== 框架构建 API ====================

@copywriting_bp.route('/projects/<project_id>/analyze-profile', methods=['POST'])
def analyze_client_profile(project_id):
    """分析客户画像"""
    logger = _get_logger()
    try:
        framework_building_agent = get_service('framework_building_agent')
        if not framework_building_agent:
            return jsonify({"success": False, "error": "框架构建器未初始化"}), 500
        
        result = framework_building_agent.analyze_client_profile(project_id)
        return jsonify(result)
        
    except Exception as e:
        logger.error(f"分析客户画像失败: {e}")
        return jsonify({"success": False, "error": str(e)}), 500


@copywriting_bp.route('/projects/<project_id>/profile-map', methods=['GET'])
def get_profile_map(project_id):
    """获取客户画像"""
    logger = _get_logger()
    try:
        framework_building_agent = get_service('framework_building_agent')
        if not framework_building_agent:
            return jsonify({"success": False, "error": "框架构建器未初始化"}), 500
        
        result = framework_building_agent.get_profile_map(project_id)
        return jsonify(result)
        
    except Exception as e:
        logger.error(f"获取客户画像失败: {e}")
        return jsonify({"success": False, "error": str(e)}), 500


@copywriting_bp.route('/projects/<project_id>/build-framework', methods=['POST'])
def build_gtv_framework(project_id):
    """构建GTV框架"""
    logger = _get_logger()
    try:
        framework_building_agent = get_service('framework_building_agent')
        content_extraction_agent = get_service('content_extraction_agent')
        
        if not framework_building_agent:
            return jsonify({"success": False, "error": "框架构建器未初始化"}), 500
        
        # 获取项目上下文
        context = ""
        if content_extraction_agent:
            context_result = content_extraction_agent.get_project_context(project_id)
            if context_result.get("success"):
                context = context_result.get("data", {}).get("context", "")
        
        # 获取客户画像（可选）
        profile_data = None
        try:
            profile_result = framework_building_agent.get_profile_map(project_id)
            if profile_result.get("success"):
                profile_data = profile_result.get("data")
        except Exception:
            pass
        
        # 构建框架
        result = framework_building_agent.build_framework(project_id, context, profile_data)
        return jsonify(result)
        
    except Exception as e:
        logger.error(f"构建GTV框架失败: {e}")
        return jsonify({"success": False, "error": str(e)}), 500


@copywriting_bp.route('/projects/<project_id>/framework', methods=['GET'])
def get_gtv_framework(project_id):
    """获取GTV框架"""
    logger = _get_logger()
    try:
        framework_building_agent = get_service('framework_building_agent')
        if not framework_building_agent:
            return jsonify({"success": False, "error": "框架构建器未初始化"}), 500
        
        result = framework_building_agent.get_framework(project_id)
        return jsonify(result)
        
    except Exception as e:
        logger.error(f"获取GTV框架失败: {e}")
        return jsonify({"success": False, "error": str(e)}), 500


@copywriting_bp.route('/projects/<project_id>/framework', methods=['PUT'])
def update_gtv_framework(project_id):
    """更新GTV框架"""
    logger = _get_logger()
    try:
        framework_building_agent = get_service('framework_building_agent')
        if not framework_building_agent:
            return jsonify({"success": False, "error": "框架构建器未初始化"}), 500
        
        data = request.get_json()
        result = framework_building_agent.update_framework(project_id, data)
        return jsonify(result)
        
    except Exception as e:
        logger.error(f"更新GTV框架失败: {e}")
        return jsonify({"success": False, "error": str(e)}), 500


@copywriting_bp.route('/projects/<project_id>/framework/export', methods=['GET'])
def export_gtv_framework(project_id):
    """导出GTV框架"""
    logger = _get_logger()
    try:
        framework_building_agent = get_service('framework_building_agent')
        if not framework_building_agent:
            return jsonify({"success": False, "error": "框架构建器未初始化"}), 500
        
        format_type = request.args.get('format', 'json')
        result = framework_building_agent.export_framework(project_id, format_type)
        
        if result.get("success") and format_type != 'json':
            file_path = result["data"]["file_path"]
            filename = result["data"]["filename"]
            return send_file(file_path, as_attachment=True, download_name=filename)
        
        return jsonify(result)
        
    except Exception as e:
        logger.error(f"导出GTV框架失败: {e}")
        return jsonify({"success": False, "error": str(e)}), 500


@copywriting_bp.route('/projects/<project_id>/framework-logs', methods=['GET'])
def get_framework_logs(project_id):
    """获取框架构建日志"""
    logger = _get_logger()
    try:
        framework_building_agent = get_service('framework_building_agent')
        if not framework_building_agent:
            return jsonify({"success": False, "error": "框架构建器未初始化"}), 500
        
        # get_framework_logs 方法只接收 project_id 参数
        result = framework_building_agent.get_framework_logs(project_id)
        
        # 返回格式化的结果
        if isinstance(result, list):
            return jsonify({"success": True, "data": result})
        return jsonify(result)
        
    except Exception as e:
        logger.error(f"获取框架日志失败: {e}")
        return jsonify({"success": False, "error": str(e)}), 500


@copywriting_bp.route('/projects/<project_id>/prompt-context', methods=['GET'])
def get_prompt_context(project_id):
    """获取提示词调试用的上下文变量"""
    logger = _get_logger()
    try:
        import sqlite3
        import sys
        import os
        # 确保模块路径正确
        project_root = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
        if project_root not in sys.path:
            sys.path.insert(0, project_root)
        
        from ace_gtv.prompts.framework_prompts import (
            MC_DESCRIPTIONS, MC_REQUIREMENTS, OC_DESCRIPTIONS, OC_REQUIREMENTS
        )
        
        db_path = _get_db_path()
        conn = sqlite3.connect(db_path)
        conn.row_factory = sqlite3.Row
        cursor = conn.cursor()
        
        # 1. 获取客户名称（从 projects 表）
        cursor.execute('''
            SELECT client_name FROM projects WHERE project_id = ? LIMIT 1
        ''', (project_id,))
        project_row = cursor.fetchone()
        client_name = project_row['client_name'] if project_row else "申请人"
        
        # 2. 获取分类后的证据摘要（从 content_classifications 表）
        cursor.execute('''
            SELECT category, subcategory, evidence_type, content, source_file, 
                   recommender_name, recommender_title, recommender_org, relationship
            FROM content_classifications 
            WHERE project_id = ? 
            ORDER BY category, created_at DESC
        ''', (project_id,))
        
        evidence_rows = cursor.fetchall()
        
        # 格式化证据文本
        evidence_items = []
        for row in evidence_rows[:20]:  # 限制数量
            content_preview = (row['content'] or '')[:300]
            if len(row['content'] or '') > 300:
                content_preview += '...'
            
            evidence_items.append(f"""【{row['category']}/{row['subcategory']}】
类型: {row['evidence_type'] or '未分类'}
内容: {content_preview}
来源: {row['source_file'] or '未知'}
""")
        
        evidence_text = "\n".join(evidence_items) if evidence_items else "暂无已分类的证据材料"
        
        # 3. 获取补充背景信息（从已提取的内容中获取）
        cursor.execute('''
            SELECT content, source_file FROM extracted_contents 
            WHERE project_id = ? 
            ORDER BY extracted_at DESC
            LIMIT 5
        ''', (project_id,))
        
        context_rows = cursor.fetchall()
        context_items = []
        for row in context_rows:
            content = row['content'] or ''
            if len(content) > 500:
                content = content[:500] + '...'
            context_items.append(f"【{row['source_file']}】\n{content}")
        
        context = "\n\n".join(context_items) if context_items else "无补充背景信息"
        
        # 4. 工作岗位选项
        work_role_options = [
            "Founder/CEO",
            "CTO/Technical Director",
            "VP/Director of Engineering",
            "Lead Engineer/Architect",
            "Product Manager/Director",
            "Research Scientist",
            "AI/ML Specialist",
            "Data Scientist/Engineer",
            "Security Specialist",
            "Investment/VC Professional",
            "Other Digital Technology Role"
        ]
        role_options = "\n".join([f"  - {role}" for role in work_role_options])
        
        # 5. 获取领域定位信息（如果已有框架）
        cursor.execute('''
            SELECT framework_data FROM gtv_frameworks 
            WHERE project_id = ? 
            ORDER BY updated_at DESC LIMIT 1
        ''', (project_id,))
        
        framework_row = cursor.fetchone()
        domain_info = ""
        framework_summary = ""
        if framework_row and framework_row['framework_data']:
            import json
            try:
                framework_data = json.loads(framework_row['framework_data'])
                domain = framework_data.get('领域定位', {})
                domain_info = f"""- 细分领域: {domain.get('细分领域', '待定')}
- 岗位定位: {domain.get('岗位定位', '待定')}
- 核心论点: {domain.get('核心论点', '待定')}
- 申请路径: {domain.get('申请路径', '待定')}"""
                
                # 框架摘要
                framework_summary = f"""领域定位: {domain.get('细分领域', '待定')}
申请路径: {domain.get('申请路径', '待定')}
核心论点: {domain.get('核心论点', '待定')}"""
            except:
                pass
        
        conn.close()
        
        # 返回所有可用变量
        variables = {
            "client_name": client_name,
            "evidence_text": evidence_text,
            "context": context[:3000] if context else "无补充信息",
            "role_options": role_options,
            "domain_info": domain_info or "待分析",
            "framework_summary": framework_summary or "待构建框架",
            # MC/OC 相关变量（默认值）
            "mc_key": "MC1_产品团队领导力",
            "mc_description": MC_DESCRIPTIONS.get("MC1_产品团队领导力", ""),
            "mc_requirement": MC_REQUIREMENTS.get("MC1_产品团队领导力", ""),
            "oc_key": "OC1_创新",
            "oc_description": OC_DESCRIPTIONS.get("OC1_创新", ""),
            "oc_requirement": OC_REQUIREMENTS.get("OC1_创新", ""),
        }
        
        return jsonify({
            "success": True,
            "data": variables
        })
        
    except Exception as e:
        logger.error(f"获取提示词上下文失败: {e}")
        import traceback
        traceback.print_exc()
        return jsonify({"success": False, "error": str(e)}), 500


# ==================== 数据清理 API ====================

@copywriting_bp.route('/projects/<project_id>/extraction/clear', methods=['POST'])
def clear_extraction_data(project_id):
    """清理项目的所有提取数据，可以重新提取"""
    logger = _get_logger()
    try:
        content_extraction_agent = get_service('content_extraction_agent')
        if not content_extraction_agent:
            return jsonify({"success": False, "error": "内容提取器未初始化"}), 500
        
        result = content_extraction_agent.clear_all_extraction_data(project_id)
        return jsonify(result)
        
    except Exception as e:
        logger.error(f"清理提取数据失败: {e}")
        return jsonify({"success": False, "error": str(e)}), 500


@copywriting_bp.route('/projects/<project_id>/framework/clear', methods=['POST'])
def clear_framework_data(project_id):
    """清理项目的所有框架数据，可以重新构建"""
    logger = _get_logger()
    try:
        framework_building_agent = get_service('framework_building_agent')
        if not framework_building_agent:
            return jsonify({"success": False, "error": "框架构建器未初始化"}), 500
        
        result = framework_building_agent.clear_all_framework_data(project_id)
        return jsonify(result)
        
    except Exception as e:
        logger.error(f"清理框架数据失败: {e}")
        return jsonify({"success": False, "error": str(e)}), 500


# ==================== 文档管理 API ====================

@copywriting_bp.route('/projects/<project_id>/documents', methods=['GET'])
def list_documents(project_id):
    """列出项目文档"""
    from pathlib import Path
    from datetime import datetime
    logger = _get_logger()
    
    project_manager = get_service('project_manager')
    if not project_manager:
        return jsonify({"success": False, "error": "服务未初始化"}), 500
    
    project = project_manager.get_project(project_id=project_id)
    if not project.get("success"):
        return jsonify(project), 404
    
    project_path = Path(project.get("path", ""))
    stage = request.args.get('stage', 'all')
    
    documents = {}
    
    stage_folders = {
        "draft": "03_文案草稿",
        "optimized": "04_优化版本",
        "final": "05_最终文档"
    }
    
    folders_to_scan = [stage_folders[stage]] if stage in stage_folders else list(stage_folders.values())
    
    for folder_name in folders_to_scan:
        folder_path = project_path / folder_name
        if folder_path.exists():
            folder_docs = {}
            for doc_type_folder in folder_path.iterdir():
                if doc_type_folder.is_dir():
                    files = []
                    for doc_file in doc_type_folder.glob("*.md"):
                        files.append({
                            "name": doc_file.name,
                            "path": str(doc_file),
                            "size": doc_file.stat().st_size,
                            "modified": datetime.fromtimestamp(doc_file.stat().st_mtime).isoformat()
                        })
                    if files:
                        folder_docs[doc_type_folder.name] = files
            if folder_docs:
                documents[folder_name] = folder_docs
    
    return jsonify({"success": True, "data": documents})


@copywriting_bp.route('/projects/<project_id>/documents/<path:doc_path>', methods=['GET'])
def get_document_content(project_id, doc_path):
    """获取文档内容"""
    from pathlib import Path
    from datetime import datetime
    logger = _get_logger()
    
    project_manager = get_service('project_manager')
    if not project_manager:
        return jsonify({"success": False, "error": "服务未初始化"}), 500
    
    project = project_manager.get_project(project_id=project_id)
    if not project.get("success"):
        return jsonify(project), 404
    
    project_path = Path(project.get("path", ""))
    file_path = project_path / doc_path
    
    if not file_path.exists():
        return jsonify({"success": False, "error": "文档不存在"}), 404
    
    try:
        with open(file_path, 'r', encoding='utf-8') as f:
            content = f.read()
        
        return jsonify({
            "success": True,
            "data": {
                "path": doc_path,
                "content": content,
                "size": file_path.stat().st_size,
                "modified": datetime.fromtimestamp(file_path.stat().st_mtime).isoformat()
            }
        })
    except Exception as e:
        logger.error(f"读取文档失败: {e}")
        return jsonify({"success": False, "error": str(e)}), 500


@copywriting_bp.route('/projects/<project_id>/documents/<path:doc_path>', methods=['PUT'])
def update_document_content(project_id, doc_path):
    """更新文档内容"""
    from pathlib import Path
    logger = _get_logger()
    
    project_manager = get_service('project_manager')
    if not project_manager:
        return jsonify({"success": False, "error": "服务未初始化"}), 500
    
    project = project_manager.get_project(project_id=project_id)
    if not project.get("success"):
        return jsonify(project), 404
    
    project_path = Path(project.get("path", ""))
    file_path = project_path / doc_path
    
    data = request.get_json()
    content = data.get('content', '')
    
    try:
        # 确保目录存在
        file_path.parent.mkdir(parents=True, exist_ok=True)
        
        with open(file_path, 'w', encoding='utf-8') as f:
            f.write(content)
        
        # 记录操作
        project_manager.log_action(project_id, "document_updated", f"更新文档: {doc_path}")
        
        return jsonify({
            "success": True,
            "message": "文档更新成功",
            "path": doc_path
        })
    except Exception as e:
        logger.error(f"更新文档失败: {e}")
        return jsonify({"success": False, "error": str(e)}), 500


# ==================== 提示词管理 API ====================

def _get_db_path():
    """获取数据库路径"""
    return DB_PATH  # 使用全局配置的数据库路径

def _ensure_system_prompts_table():
    """确保系统提示词表存在并初始化默认数据"""
    import sqlite3
    logger = _get_logger()
    db_path = _get_db_path()
    
    try:
        conn = sqlite3.connect(db_path)
        cursor = conn.cursor()
        
        cursor.execute('''
            CREATE TABLE IF NOT EXISTS system_prompts (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                name TEXT NOT NULL UNIQUE,
                type TEXT NOT NULL,
                description TEXT,
                content TEXT NOT NULL,
                version INTEGER DEFAULT 1,
                is_active INTEGER DEFAULT 1,
                created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
                updated_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
            )
        ''')
        
        # 检查是否需要添加 version 列
        cursor.execute("PRAGMA table_info(system_prompts)")
        columns = [row[1] for row in cursor.fetchall()]
        if 'version' not in columns:
            cursor.execute('ALTER TABLE system_prompts ADD COLUMN version INTEGER DEFAULT 1')
        
        # 默认提示词
        default_prompts = [
            {
                "name": "框架构建提示词",
                "type": "framework",
                "description": "用于构建GTV申请框架的主提示词",
                "content": """你是一位专业的GTV签证申请顾问，请根据以下材料为申请人构建申请框架。

## 申请人信息
{client_info}

## 材料内容
{context}

请按照MC/OC标准分析材料，输出JSON格式的框架。"""
            },
            {
                "name": "内容提取提示词",
                "type": "extraction",
                "description": "用于从材料中提取关键内容",
                "content": """请从以下材料中提取GTV申请相关的关键信息。

## 材料内容
{content}

请识别并提取：
1. 申请人的成就和贡献
2. 行业认可和奖项
3. 专业技能和经验
4. 推荐人信息"""
            },
            {
                "name": "英文翻译提示词",
                "type": "translation",
                "description": "将英文内容翻译为中文",
                "content": """请将以下英文内容翻译成中文。

## 英文原文
{content}

请返回翻译后的中文内容。"""
            }
        ]
        
        # 检查并插入默认提示词
        cursor.execute("SELECT COUNT(*) FROM system_prompts")
        if cursor.fetchone()[0] == 0:
            for p in default_prompts:
                cursor.execute('''
                    INSERT OR IGNORE INTO system_prompts (name, type, description, content, version)
                    VALUES (?, ?, ?, ?, 1)
                ''', (p['name'], p['type'], p['description'], p['content']))
        
        conn.commit()
        conn.close()
        logger.info("系统提示词表初始化完成")
    except Exception as e:
        logger.error(f"初始化系统提示词表失败: {e}")


@copywriting_bp.route('/agent-prompts', methods=['GET'])
def get_system_prompts():
    """获取所有系统提示词"""
    import sqlite3
    logger = _get_logger()
    db_path = _get_db_path()
    
    try:
        _ensure_system_prompts_table()
        
        conn = sqlite3.connect(db_path)
        conn.row_factory = sqlite3.Row
        cursor = conn.cursor()
        
        cursor.execute('''
            SELECT id, name, type, description, content, version, is_active, created_at, updated_at
            FROM system_prompts
            ORDER BY type, name
        ''')
        
        prompts = []
        for row in cursor.fetchall():
            prompts.append({
                "id": row['id'],
                "name": row['name'],
                "type": row['type'],
                "description": row['description'],
                "content": row['content'],
                "version": row['version'] or 1,
                "is_active": bool(row['is_active']),
                "created_at": row['created_at'],
                "updated_at": row['updated_at']
            })
        
        conn.close()
        return jsonify({"success": True, "data": prompts})
        
    except Exception as e:
        logger.error(f"获取系统提示词失败: {e}")
        return jsonify({"success": False, "error": str(e)}), 500


@copywriting_bp.route('/agent-prompts/<int:prompt_id>', methods=['GET'])
def get_system_prompt(prompt_id):
    """获取单个系统提示词"""
    import sqlite3
    logger = _get_logger()
    db_path = _get_db_path()
    
    try:
        conn = sqlite3.connect(db_path)
        conn.row_factory = sqlite3.Row
        cursor = conn.cursor()
        
        cursor.execute('''
            SELECT id, name, type, description, content, version, is_active, created_at, updated_at
            FROM system_prompts WHERE id = ?
        ''', (prompt_id,))
        
        row = cursor.fetchone()
        conn.close()
        
        if not row:
            return jsonify({"success": False, "error": "提示词不存在"}), 404
        
        return jsonify({
            "success": True,
            "data": {
                "id": row['id'],
                "name": row['name'],
                "type": row['type'],
                "description": row['description'],
                "content": row['content'],
                "version": row['version'] or 1,
                "is_active": bool(row['is_active']),
                "created_at": row['created_at'],
                "updated_at": row['updated_at']
            }
        })
        
    except Exception as e:
        logger.error(f"获取提示词失败: {e}")
        return jsonify({"success": False, "error": str(e)}), 500


@copywriting_bp.route('/agent-prompts/<int:prompt_id>', methods=['PUT'])
def update_system_prompt(prompt_id):
    """更新系统提示词（自动增加版本号并保存历史）"""
    import sqlite3
    logger = _get_logger()
    db_path = _get_db_path()
    
    try:
        data = request.get_json()
        
        conn = sqlite3.connect(db_path)
        conn.row_factory = sqlite3.Row
        cursor = conn.cursor()
        
        # 获取当前提示词信息
        cursor.execute('SELECT name, version, content FROM system_prompts WHERE id = ?', (prompt_id,))
        row = cursor.fetchone()
        if not row:
            return jsonify({"success": False, "error": "提示词不存在"}), 404
        
        current_name = row['name']
        current_version = (row['version'] or 1) if row else 1
        current_content = row['content']
        new_content = data.get('content', current_content)
        
        # 只有内容有变化才增加版本号
        if new_content != current_content:
            new_version = current_version + 1
        else:
            new_version = current_version
        
        # 更新主表
        cursor.execute('''
            UPDATE system_prompts 
            SET name = ?, description = ?, content = ?, is_active = ?, 
                version = ?, updated_at = CURRENT_TIMESTAMP
            WHERE id = ?
        ''', (
            data.get('name', current_name),
            data.get('description'),
            new_content,
            1 if data.get('is_active', True) else 0,
            new_version,
            prompt_id
        ))
        
        # 如果内容有变化，保存到历史表
        if new_content != current_content:
            # 确保历史表存在
            cursor.execute('''
                CREATE TABLE IF NOT EXISTS system_prompts_history (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    prompt_name TEXT NOT NULL,
                    version INTEGER NOT NULL,
                    content TEXT NOT NULL,
                    description TEXT,
                    changed_by TEXT,
                    change_reason TEXT,
                    created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
                    UNIQUE(prompt_name, version)
                )
            ''')
            
            cursor.execute('''
                INSERT OR REPLACE INTO system_prompts_history 
                (prompt_name, version, content, description, changed_by, change_reason)
                VALUES (?, ?, ?, ?, ?, ?)
            ''', (
                data.get('name', current_name), 
                new_version, 
                new_content, 
                data.get('description'),
                "用户",
                "手动编辑"
            ))
            
            logger.info(f"提示词 {prompt_id} ({current_name}) 已更新至版本 v{new_version}")
        
        conn.commit()
        conn.close()
        
        return jsonify({
            "success": True, 
            "message": f"提示词已更新至版本 v{new_version}" if new_content != current_content else "提示词已保存",
            "version": new_version
        })
        
    except Exception as e:
        logger.error(f"更新提示词失败: {e}")
        import traceback
        traceback.print_exc()
        return jsonify({"success": False, "error": str(e)}), 500


@copywriting_bp.route('/agent-prompts/<int:prompt_id>/history', methods=['GET'])
def get_prompt_history(prompt_id):
    """获取提示词的版本历史"""
    import sqlite3
    logger = _get_logger()
    db_path = _get_db_path()
    
    try:
        conn = sqlite3.connect(db_path)
        conn.row_factory = sqlite3.Row
        cursor = conn.cursor()
        
        # 首先获取提示词名称
        cursor.execute('SELECT name FROM system_prompts WHERE id = ?', (prompt_id,))
        row = cursor.fetchone()
        if not row:
            return jsonify({"success": False, "error": "提示词不存在"}), 404
        
        prompt_name = row['name']
        
        # 获取版本历史
        cursor.execute('''
            SELECT version, content, description, changed_by, change_reason, created_at
            FROM system_prompts_history 
            WHERE prompt_name = ?
            ORDER BY version DESC
        ''', (prompt_name,))
        
        history = []
        for row in cursor.fetchall():
            history.append({
                "version": row['version'],
                "content": row['content'],
                "description": row['description'],
                "changed_by": row['changed_by'],
                "change_reason": row['change_reason'],
                "created_at": row['created_at']
            })
        
        conn.close()
        
        return jsonify({
            "success": True,
            "data": {
                "prompt_name": prompt_name,
                "history": history
            }
        })
        
    except Exception as e:
        logger.error(f"获取提示词历史失败: {e}")
        return jsonify({"success": False, "error": str(e)}), 500


@copywriting_bp.route('/agent-prompts', methods=['POST'])
def create_system_prompt():
    """创建新的系统提示词"""
    import sqlite3
    logger = _get_logger()
    db_path = _get_db_path()
    
    try:
        data = request.get_json()
        
        conn = sqlite3.connect(db_path)
        cursor = conn.cursor()
        
        cursor.execute('''
            INSERT INTO system_prompts (name, type, description, content)
            VALUES (?, ?, ?, ?)
        ''', (
            data.get('name'),
            data.get('type', 'custom'),
            data.get('description', ''),
            data.get('content', '')
        ))
        
        prompt_id = cursor.lastrowid
        conn.commit()
        conn.close()
        
        return jsonify({"success": True, "data": {"id": prompt_id}})
        
    except Exception as e:
        logger.error(f"创建提示词失败: {e}")
        return jsonify({"success": False, "error": str(e)}), 500


@copywriting_bp.route('/agent-prompts/sync', methods=['POST'])
def sync_system_prompts():
    """同步系统默认提示词（从代码模板同步到数据库）"""
    logger = _get_logger()
    
    try:
        _ensure_system_prompts_table()
        
        import sqlite3
        import sys
        import os
        # 确保模块路径正确
        project_root = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
        if project_root not in sys.path:
            sys.path.insert(0, project_root)
        
        from ace_gtv.prompts import FRAMEWORK_PROMPTS_CONFIG
        
        db_path = _get_db_path()
        conn = sqlite3.connect(db_path)
        cursor = conn.cursor()
        
        # 确保版本历史表存在
        cursor.execute('''
            CREATE TABLE IF NOT EXISTS system_prompts_history (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                prompt_name TEXT NOT NULL,
                version INTEGER NOT NULL,
                content TEXT NOT NULL,
                description TEXT,
                changed_by TEXT,
                change_reason TEXT,
                created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
                UNIQUE(prompt_name, version)
            )
        ''')
        
        updated_count = 0
        inserted_count = 0
        
        for prompt_config in FRAMEWORK_PROMPTS_CONFIG:
            prompt_type = prompt_config['type']
            prompt_name = prompt_config['name']
            prompt_content = prompt_config['content']
            prompt_description = prompt_config['description']
            prompt_category = prompt_config.get('category', 'framework')
            
            # 检查是否已存在
            cursor.execute(
                'SELECT id, version, content FROM system_prompts WHERE type = ?',
                (prompt_type,)
            )
            existing = cursor.fetchone()
            
            if existing:
                existing_id, existing_version, existing_content = existing
                # 只有内容有变化才更新
                if existing_content != prompt_content:
                    new_version = (existing_version or 0) + 1
                    cursor.execute('''
                        UPDATE system_prompts 
                        SET content = ?, description = ?, version = ?, updated_at = CURRENT_TIMESTAMP
                        WHERE id = ?
                    ''', (prompt_content, prompt_description, new_version, existing_id))
                    
                    # 保存到历史表
                    cursor.execute('''
                        INSERT OR REPLACE INTO system_prompts_history 
                        (prompt_name, version, content, description, changed_by, change_reason)
                        VALUES (?, ?, ?, ?, ?, ?)
                    ''', (prompt_name, new_version, prompt_content, prompt_description, 
                          "系统同步", "从代码模板同步"))
                    
                    updated_count += 1
                    logger.info(f"更新提示词: {prompt_name} -> v{new_version}")
            else:
                # 插入新提示词
                cursor.execute('''
                    INSERT INTO system_prompts (name, type, description, content, category, version)
                    VALUES (?, ?, ?, ?, ?, 1)
                ''', (prompt_name, prompt_type, prompt_description, prompt_content, prompt_category))
                
                # 保存到历史表
                cursor.execute('''
                    INSERT OR REPLACE INTO system_prompts_history 
                    (prompt_name, version, content, description, changed_by, change_reason)
                    VALUES (?, ?, ?, ?, ?, ?)
                ''', (prompt_name, 1, prompt_content, prompt_description, "系统同步", "初始版本"))
                
                inserted_count += 1
                logger.info(f"插入提示词: {prompt_name} v1")
        
        conn.commit()
        
        # 获取总数
        cursor.execute('SELECT COUNT(*) FROM system_prompts')
        total_count = cursor.fetchone()[0]
        
        conn.close()
        
        return jsonify({
            "success": True, 
            "message": f"提示词同步完成：新增 {inserted_count} 个，更新 {updated_count} 个，共 {total_count} 个",
            "count": total_count,
            "inserted": inserted_count,
            "updated": updated_count
        })
        
    except Exception as e:
        logger.error(f"同步提示词失败: {e}")
        import traceback
        traceback.print_exc()
        return jsonify({"success": False, "error": str(e)}), 500


@copywriting_bp.route('/agent-prompts/debug', methods=['POST'])
def debug_prompt():
    """调试提示词"""
    import re
    import os
    import requests
    logger = _get_logger()
    
    try:
        data = request.get_json()
        prompt_content = data.get('prompt_content', '')
        variables = data.get('variables', {})
        
        if not prompt_content:
            return jsonify({"success": False, "error": "提示词内容为空"}), 400
        
        # 替换变量
        final_prompt = prompt_content
        for key, value in variables.items():
            final_prompt = final_prompt.replace(f"{{{key}}}", str(value))
        
        # 检查未替换的变量
        remaining_vars = re.findall(r'\{(\w+)\}', final_prompt)
        if remaining_vars:
            for var in remaining_vars:
                final_prompt = final_prompt.replace(f"{{{var}}}", f"[未提供: {var}]")
        
        # 根据 LLM_PROVIDER 配置选择调用方式
        llm_provider = os.getenv("LLM_PROVIDER", "ENNCLOUD").upper()
        
        if llm_provider == "ENNCLOUD":
            # 使用 EnnCloud API
            api_key = os.getenv("ENNCLOUD_API_KEY")
            base_url = os.getenv("ENNCLOUD_BASE_URL", "https://ai.enncloud.cn/v1")
            model = os.getenv("ENNCLOUD_MODEL", "GLM-4.5-Air")
            
            if not api_key:
                return jsonify({"success": False, "error": "ENNCLOUD_API_KEY 未配置"}), 500
            
            url = f"{base_url}/chat/completions"
            headers = {
                "Content-Type": "application/json",
                "Authorization": f"Bearer {api_key}"
            }
            payload = {
                "model": model,
                "messages": [{"role": "user", "content": final_prompt}],
                "temperature": 0.7,
                "max_tokens": 4000,
                "stream": False
            }
            response = requests.post(url, headers=headers, json=payload, timeout=120)
            response.raise_for_status()
            resp_data = response.json()
            output = resp_data["choices"][0]["message"]["content"]
            tokens_used = resp_data.get("usage", {}).get("total_tokens", 0)
            
        elif llm_provider == "ANTHROPIC":
            # 使用 Anthropic API
            from anthropic import Anthropic
            api_key = os.getenv("ANTHROPIC_API_KEY")
            base_url = os.getenv("ANTHROPIC_BASE_URL")
            
            if not api_key:
                return jsonify({"success": False, "error": "ANTHROPIC_API_KEY 未配置"}), 500
            
            if base_url:
                client = Anthropic(api_key=api_key, base_url=base_url)
            else:
                client = Anthropic(api_key=api_key)
            
            model = os.getenv("ANTHROPIC_MODEL", "claude-3-5-sonnet-20241022")
            response = client.messages.create(
                model=model,
                max_tokens=4000,
                messages=[{"role": "user", "content": final_prompt}]
            )
            output = response.content[0].text if response.content else ""
            tokens_used = response.usage.input_tokens + response.usage.output_tokens
            
        else:
            return jsonify({"success": False, "error": f"不支持的 LLM 提供商: {llm_provider}"}), 500
        
        return jsonify({
            "success": True,
            "data": {
                "input": final_prompt[:500] + "..." if len(final_prompt) > 500 else final_prompt,
                "output": output,
                "tokens_used": tokens_used
            }
        })
        
    except requests.exceptions.RequestException as e:
        logger.error(f"调试提示词失败 (网络错误): {e}")
        return jsonify({"success": False, "error": f"网络请求失败: {str(e)}"}), 500
    except Exception as e:
        logger.error(f"调试提示词失败: {e}")
        return jsonify({"success": False, "error": str(e)}), 500


@copywriting_bp.route('/projects/<project_id>/material-collection/download-all', methods=['GET'])
def download_all_materials(project_id):
    """
    打包下载项目所有材料
    按分类建立文件夹，将文件放入对应分类，打包成ZIP下载
    """
    import zipfile
    import tempfile
    import shutil
    import re
    from datetime import datetime
    
    logger = _get_logger()
    _init_services()
    raw_material_manager = _services.get('raw_material_manager')
    db = _services.get('db')
    
    if not raw_material_manager:
        return jsonify({"success": False, "error": "服务未初始化"}), 500
    
    try:
        # 获取项目信息
        project_result = db.get_project(project_id=project_id) if db else None
        client_name = "unknown"
        if project_result and project_result.get("success"):
            client_name = project_result.get("data", {}).get("client_name", project_id)
        
        # 安全处理客户名称（保留中文、字母、数字、空格、连字符、下划线）
        safe_client_name = re.sub(r'[^\w\s\-]', '', client_name, flags=re.UNICODE).strip()
        safe_client_name = safe_client_name.replace('/', '_').replace('\\', '_')
        if not safe_client_name:
            safe_client_name = project_id
        
        # 获取材料分类结构
        from services.raw_material_manager import MATERIAL_CATEGORIES
        
        # 从数据库获取项目的所有材料文件
        conn = sqlite3.connect(raw_material_manager.db_path)
        conn.row_factory = sqlite3.Row
        cursor = conn.cursor()
        
        cursor.execute('''
            SELECT mf.*, mc.item_id as mc_item_id
            FROM material_files mf
            LEFT JOIN material_collection mc 
            ON mf.project_id = mc.project_id 
            AND mf.category_id = mc.category_id 
            AND mf.item_id = mc.item_id
            WHERE mf.project_id = ?
        ''', (project_id,))
        
        files = cursor.fetchall()
        conn.close()
        
        if not files:
            return jsonify({"success": False, "error": "该项目暂无材料文件"}), 404
        
        # 创建临时目录
        temp_dir = tempfile.mkdtemp()
        zip_base_dir = os.path.join(temp_dir, f"{safe_client_name}_材料")
        os.makedirs(zip_base_dir, exist_ok=True)
        
        files_copied = 0
        errors = []
        
        # 遍历文件并按分类组织
        for file_row in files:
            category_id = file_row['category_id']
            item_id = file_row['item_id']
            file_path = file_row['file_path']
            file_name = file_row['file_name']
            
            # 获取标签名称
            folder_name = "其他材料"
            
            if category_id in MATERIAL_CATEGORIES:
                cat = MATERIAL_CATEGORIES[category_id]
                category_name = cat.get('name', category_id)
                for item in cat.get('items', []):
                    if item.get('item_id') == item_id:
                        folder_name = item.get('name', category_name)
                        break
                else:
                    folder_name = category_name
            
            # 创建文件夹
            folder_path = os.path.join(zip_base_dir, folder_name)
            os.makedirs(folder_path, exist_ok=True)
            
            # 复制文件
            if os.path.exists(file_path):
                dest_path = os.path.join(folder_path, file_name)
                if os.path.exists(dest_path):
                    base, ext = os.path.splitext(file_name)
                    counter = 1
                    while os.path.exists(dest_path):
                        dest_path = os.path.join(folder_path, f"{base}_{counter}{ext}")
                        counter += 1
                
                shutil.copy2(file_path, dest_path)
                files_copied += 1
            else:
                errors.append(f"文件不存在: {file_name}")
        
        if files_copied == 0:
            shutil.rmtree(temp_dir)
            return jsonify({"success": False, "error": "没有可下载的文件"}), 404
        
        # 创建ZIP文件
        zip_filename = f"{safe_client_name}_材料_{datetime.now().strftime('%Y%m%d_%H%M%S')}.zip"
        zip_path = os.path.join(temp_dir, zip_filename)
        
        with zipfile.ZipFile(zip_path, 'w', zipfile.ZIP_DEFLATED) as zipf:
            for root, dirs, files_in_dir in os.walk(zip_base_dir):
                for file in files_in_dir:
                    file_full_path = os.path.join(root, file)
                    arcname = os.path.relpath(file_full_path, zip_base_dir)
                    zipf.write(file_full_path, arcname)
        
        logger.info(f"材料打包完成: {zip_path}, 共 {files_copied} 个文件")
        
        # 发送文件
        response = send_file(
            zip_path,
            mimetype='application/zip',
            as_attachment=True,
            download_name=zip_filename
        )
        
        # 注册清理回调
        @response.call_on_close
        def cleanup():
            try:
                shutil.rmtree(temp_dir)
            except Exception as e:
                logger.warning(f"清理临时目录失败: {e}")
        
        return response
        
    except Exception as e:
        logger.error(f"打包下载材料失败: {e}")
        import traceback
        traceback.print_exc()
        return jsonify({"success": False, "error": str(e)}), 500


# ==================== 文案生成 - 材料包管理 API ====================

@copywriting_bp.route('/projects/<project_id>/generation-context', methods=['GET'])
def get_generation_context(project_id):
    """
    获取文案生成的聚合上下文
    包括提取的分类内容、原始材料摘要、GTV框架信息
    """
    import sqlite3
    logger = _get_logger()
    
    try:
        _init_services()
        db = _services.get('db')
        
        if not db:
            return jsonify({"success": False, "error": "服务未初始化"}), 500
        
        context = {
            "project_id": project_id,
            "classifications": [],
            "raw_materials": [],
            "framework": None,
            "recommenders": []
        }
        
        # 获取提取的分类内容
        conn = sqlite3.connect(db.db_path)
        conn.row_factory = sqlite3.Row
        cursor = conn.cursor()
        
        cursor.execute('''
            SELECT id, category, subcategory, content, source_file, 
                   relevance_score, key_points, recommender_name, 
                   recommender_title, recommender_org, relationship
            FROM content_classifications
            WHERE project_id = ?
            ORDER BY category, relevance_score DESC
        ''', (project_id,))
        
        classifications = []
        for row in cursor.fetchall():
            classifications.append({
                "id": row['id'],
                "category": row['category'],
                "subcategory": row['subcategory'],
                "content": row['content'],
                "source_file": row['source_file'],
                "relevance_score": row['relevance_score'],
                "key_points": json.loads(row['key_points']) if row['key_points'] else [],
                "recommender_name": row['recommender_name'],
                "recommender_title": row['recommender_title'],
                "recommender_org": row['recommender_org'],
                "relationship": row['relationship']
            })
        context["classifications"] = classifications
        
        # 获取原始材料列表
        cursor.execute('''
            SELECT id, file_name, file_type, category_id, file_size, upload_time
            FROM material_files
            WHERE project_id = ?
        ''', (project_id,))
        
        materials = []
        for row in cursor.fetchall():
            materials.append({
                "id": row['id'],
                "file_name": row['file_name'],
                "file_type": row['file_type'],
                "category": row['category_id'],
                "file_size": row['file_size'],
                "upload_time": row['upload_time']
            })
        context["raw_materials"] = materials
        
        # 获取GTV框架
        cursor.execute('''
            SELECT framework_data FROM gtv_framework
            WHERE project_id = ?
        ''', (project_id,))
        
        row = cursor.fetchone()
        if row and row['framework_data']:
            context["framework"] = json.loads(row['framework_data'])
            
            # 提取推荐人信息
            if context["framework"] and "推荐信" in context["framework"]:
                rec_info = context["framework"]["推荐信"]
                for key in ["推荐人1", "推荐人2", "推荐人3"]:
                    if key in rec_info and rec_info[key]:
                        context["recommenders"].append({
                            "key": key,
                            **rec_info[key]
                        })
        
        conn.close()
        
        return jsonify({
            "success": True,
            "data": context
        })
        
    except Exception as e:
        logger.error(f"获取生成上下文失败: {e}")
        import traceback
        traceback.print_exc()
        return jsonify({"success": False, "error": str(e)}), 500


@copywriting_bp.route('/projects/<project_id>/packages/<package_type>', methods=['GET'])
def get_package_content(project_id, package_type):
    """获取材料包内容"""
    import sqlite3
    logger = _get_logger()
    
    try:
        _init_services()
        db = _services.get('db')
        
        if not db:
            return jsonify({"success": False, "error": "服务未初始化"}), 500
        
        conn = sqlite3.connect(db.db_path)
        conn.row_factory = sqlite3.Row
        cursor = conn.cursor()
        
        # 确保表存在
        cursor.execute('''
            CREATE TABLE IF NOT EXISTS package_contents (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                project_id TEXT NOT NULL,
                package_type TEXT NOT NULL,
                current_version INTEGER DEFAULT 1,
                content TEXT,
                status TEXT DEFAULT 'draft',
                last_edited_by TEXT,
                ai_generated INTEGER DEFAULT 0,
                created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
                updated_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
                UNIQUE(project_id, package_type)
            )
        ''')
        conn.commit()
        
        cursor.execute('''
            SELECT * FROM package_contents
            WHERE project_id = ? AND package_type = ?
        ''', (project_id, package_type))
        
        row = cursor.fetchone()
        conn.close()
        
        if row:
            return jsonify({
                "success": True,
                "data": {
                    "project_id": row['project_id'],
                    "package_type": row['package_type'],
                    "current_version": row['current_version'],
                    "content": row['content'],
                    "status": row['status'],
                    "last_edited_by": row['last_edited_by'],
                    "ai_generated": bool(row['ai_generated']),
                    "updated_at": row['updated_at']
                }
            })
        else:
            return jsonify({
                "success": True,
                "data": {
                    "project_id": project_id,
                    "package_type": package_type,
                    "current_version": 0,
                    "content": "",
                    "status": "draft",
                    "last_edited_by": None,
                    "ai_generated": False,
                    "updated_at": None
                }
            })
        
    except Exception as e:
        logger.error(f"获取材料包内容失败: {e}")
        return jsonify({"success": False, "error": str(e)}), 500


@copywriting_bp.route('/projects/<project_id>/packages/<package_type>', methods=['POST'])
def save_package_content(project_id, package_type):
    """保存材料包内容"""
    import sqlite3
    logger = _get_logger()
    
    try:
        _init_services()
        db = _services.get('db')
        
        if not db:
            return jsonify({"success": False, "error": "服务未初始化"}), 500
        
        data = request.get_json()
        content = data.get('content', '')
        edit_type = data.get('edit_type', 'manual')
        edit_summary = data.get('edit_summary', '保存内容')
        
        conn = sqlite3.connect(db.db_path)
        conn.row_factory = sqlite3.Row
        cursor = conn.cursor()
        
        # 确保表存在
        cursor.execute('''
            CREATE TABLE IF NOT EXISTS package_contents (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                project_id TEXT NOT NULL,
                package_type TEXT NOT NULL,
                current_version INTEGER DEFAULT 1,
                content TEXT,
                status TEXT DEFAULT 'draft',
                last_edited_by TEXT,
                ai_generated INTEGER DEFAULT 0,
                created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
                updated_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
                UNIQUE(project_id, package_type)
            )
        ''')
        
        cursor.execute('''
            CREATE TABLE IF NOT EXISTS document_versions (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                project_id TEXT NOT NULL,
                package_type TEXT NOT NULL,
                version INTEGER NOT NULL,
                content TEXT,
                edit_type TEXT,
                edit_summary TEXT,
                editor TEXT,
                word_count INTEGER,
                created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
            )
        ''')
        conn.commit()
        
        # 获取当前版本号
        cursor.execute('''
            SELECT current_version FROM package_contents
            WHERE project_id = ? AND package_type = ?
        ''', (project_id, package_type))
        
        row = cursor.fetchone()
        new_version = (row['current_version'] + 1) if row else 1
        
        # 保存版本历史
        word_count = len(content.split()) if content else 0
        cursor.execute('''
            INSERT INTO document_versions 
            (project_id, package_type, version, content, edit_type, edit_summary, editor, word_count)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?)
        ''', (project_id, package_type, new_version, content, edit_type, edit_summary, 'user', word_count))
        
        # 更新或插入当前内容
        cursor.execute('''
            INSERT INTO package_contents 
            (project_id, package_type, current_version, content, status, last_edited_by, ai_generated, updated_at)
            VALUES (?, ?, ?, ?, 'draft', 'user', 0, CURRENT_TIMESTAMP)
            ON CONFLICT(project_id, package_type) 
            DO UPDATE SET 
                current_version = excluded.current_version,
                content = excluded.content,
                status = 'draft',
                last_edited_by = 'user',
                updated_at = CURRENT_TIMESTAMP
        ''', (project_id, package_type, new_version, content))
        
        conn.commit()
        conn.close()
        
        logger.info(f"保存材料包内容: {project_id}/{package_type} v{new_version}")
        
        return jsonify({
            "success": True,
            "version": new_version,
            "message": f"已保存为版本 {new_version}"
        })
        
    except Exception as e:
        logger.error(f"保存材料包内容失败: {e}")
        return jsonify({"success": False, "error": str(e)}), 500


@copywriting_bp.route('/projects/<project_id>/packages/<package_type>/generate', methods=['POST'])
def generate_package_content(project_id, package_type):
    """使用AI生成材料包内容"""
    import sqlite3
    import os
    import requests
    logger = _get_logger()
    
    try:
        _init_services()
        db = _services.get('db')
        
        if not db:
            return jsonify({"success": False, "error": "服务未初始化"}), 500
        
        data = request.get_json() or {}
        custom_instructions = data.get('custom_instructions', '')
        selected_inputs = data.get('selected_inputs', {})
        recommender_info = data.get('recommender_info', {})
        custom_prompt = data.get('custom_prompt', '')  # 用户自定义的系统提示词
        include_strategy = data.get('include_strategy', True)  # 是否包含GTV策略信息
        
        # 获取生成上下文
        conn = sqlite3.connect(db.db_path)
        conn.row_factory = sqlite3.Row
        cursor = conn.cursor()
        
        # 构建上下文
        context_parts = []
        
        # 获取项目信息
        cursor.execute('SELECT * FROM projects WHERE project_id = ?', (project_id,))
        project_row = cursor.fetchone()
        if project_row:
            context_parts.append(f"申请人: {project_row['client_name']}")
            context_parts.append(f"签证类型: {project_row['visa_type']}")
        
        # 获取GTV框架（增强版：包含更多策略信息）
        include_framework = selected_inputs.get('include_framework', True) and include_strategy
        if include_framework:
            cursor.execute('SELECT framework_data, mc_criteria, oc_criteria, domain_positioning FROM gtv_framework WHERE project_id = ?', (project_id,))
            fw_row = cursor.fetchone()
            if fw_row:
                context_parts.append("\n--- GTV申请策略 ---")
                
                # 解析框架数据
                if fw_row['framework_data']:
                    framework = json.loads(fw_row['framework_data'])
                    
                    # 领域定位
                    if framework.get('领域定位'):
                        domain = framework['领域定位']
                        context_parts.append(f"\n【领域定位】")
                        if isinstance(domain, dict):
                            if domain.get('核心领域'):
                                context_parts.append(f"核心领域: {domain['核心领域']}")
                            if domain.get('岗位定位'):
                                context_parts.append(f"岗位定位: {domain['岗位定位']}")
                            if domain.get('核心论点'):
                                context_parts.append(f"核心论点: {domain['核心论点']}")
                            if domain.get('关键成就'):
                                context_parts.append(f"关键成就: {domain['关键成就']}")
                        else:
                            context_parts.append(str(domain))
                    
                    # MC标准详情
                    if framework.get('MC_必选标准'):
                        mc = framework['MC_必选标准']
                        context_parts.append(f"\n【MC标准 - 必选】")
                        context_parts.append(f"选择的MC: {mc.get('选择的MC', '')}")
                        if mc.get('核心论点'):
                            context_parts.append(f"核心论点: {mc['核心论点']}")
                        if mc.get('支撑证据'):
                            context_parts.append(f"支撑证据: {mc['支撑证据']}")
                    
                    # OC标准详情
                    if framework.get('OC_可选标准'):
                        oc = framework['OC_可选标准']
                        context_parts.append(f"\n【OC标准 - 可选】")
                        selected_ocs = oc.get('选择的OC', [])
                        context_parts.append(f"选择的OC: {', '.join(selected_ocs) if isinstance(selected_ocs, list) else selected_ocs}")
                        for oc_key in ['OC1', 'OC2', 'OC3', 'OC4']:
                            if oc.get(oc_key):
                                oc_detail = oc[oc_key]
                                if isinstance(oc_detail, dict) and oc_detail.get('核心论点'):
                                    context_parts.append(f"{oc_key}: {oc_detail['核心论点']}")
                    
                    # 推荐信策略（如果是推荐信类型）
                    if package_type.startswith('rl_') and framework.get('推荐信'):
                        rl_strategy = framework['推荐信']
                        context_parts.append(f"\n【推荐信策略】")
                        for key, value in rl_strategy.items():
                            if value and isinstance(value, dict):
                                context_parts.append(f"{key}: {json.dumps(value, ensure_ascii=False)}")
        
        # 获取分类内容
        classification_ids = selected_inputs.get('classification_ids', [])
        if classification_ids:
            placeholders = ','.join('?' * len(classification_ids))
            cursor.execute(f'''
                SELECT category, subcategory, content FROM content_classifications
                WHERE project_id = ? AND id IN ({placeholders})
            ''', [project_id] + classification_ids)
            
            context_parts.append("\n--- 提取的内容 ---")
            for row in cursor.fetchall():
                context_parts.append(f"\n[{row['category']}] {row['subcategory'] or ''}:")
                context_parts.append(row['content'][:2000])
        
        conn.close()
        
        # 添加推荐人信息（如果是推荐信）
        if package_type.startswith('rl_') and recommender_info:
            context_parts.append("\n--- 推荐人信息 ---")
            context_parts.append(f"姓名: {recommender_info.get('name', '待填写')}")
            context_parts.append(f"职位: {recommender_info.get('title', '待填写')}")
            context_parts.append(f"机构: {recommender_info.get('organization', '待填写')}")
            context_parts.append(f"关系: {recommender_info.get('relationship', '待填写')}")
        
        # 添加自定义指令
        if custom_instructions:
            context_parts.append(f"\n--- 特别要求 ---\n{custom_instructions}")
        
        context = "\n".join(context_parts)
        
        # 获取材料包类型的系统提示词
        package_prompts = {
            "personal_statement": "你是一位专业的GTV签证个人陈述撰写专家。请根据提供的信息，撰写一份有说服力的个人陈述（英文），1000-1500词。",
            "cv_resume": "你是一位专业的简历优化专家。请根据提供的信息，生成一份符合GTV签证要求的专业简历（英文）。",
            "recommendation_letters": "你是一位专业的推荐信撰写专家。请根据提供的信息，撰写一封GTV签证推荐信（英文），1-1.5页。推荐信应突出申请人在相关领域的卓越成就和贡献。",
            "rl_1": "你是一位专业的推荐信撰写专家。请撰写一封来自行业专家的GTV签证推荐信（英文），1-1.5页。",
            "rl_2": "你是一位专业的推荐信撰写专家。请撰写一封来自技术/学术专家的GTV签证推荐信（英文），1-1.5页。",
            "rl_3": "你是一位专业的推荐信撰写专家。请撰写一封来自商业合作伙伴的GTV签证推荐信（英文），1-1.5页。",
            "evidence_portfolio": "你是一位GTV签证证据材料整理专家。请根据提供的信息，整理证据材料清单和说明（中英双语），包括每项证据的重要性和支撑论点。",
            "cover_letter": "你是一位专业的签证申请信撰写专家。请撰写一封正式的GTV签证申请信（英文），简洁有力地概述申请人的资格和申请理由。",
            "endorsement_letter": "你是一位Tech Nation背书申请专家。请撰写背书申请材料。",
            "business_plan": "你是一位商业计划书撰写专家。请撰写创业者路径所需的商业计划书。",
            "supplementary": "你是一位签证材料整理专家。请整理补充材料说明。"
        }
        
        # 使用自定义提示词或默认提示词
        if custom_prompt and custom_prompt.strip():
            system_prompt = custom_prompt
        else:
            system_prompt = package_prompts.get(package_type, "你是一位专业的签证材料撰写专家。")
        
        # 调用LLM
        llm_provider = os.getenv("LLM_PROVIDER", "ENNCLOUD").upper()
        
        if llm_provider == "ENNCLOUD":
            api_key = os.getenv("ENNCLOUD_API_KEY")
            base_url = os.getenv("ENNCLOUD_BASE_URL", "https://ai.enncloud.cn/v1")
            model = os.getenv("ENNCLOUD_MODEL", "GLM-4.5-Air")
            
            if not api_key:
                return jsonify({"success": False, "error": "ENNCLOUD_API_KEY 未配置"}), 500
            
            url = f"{base_url}/chat/completions"
            headers = {
                "Content-Type": "application/json",
                "Authorization": f"Bearer {api_key}"
            }
            payload = {
                "model": model,
                "messages": [
                    {"role": "system", "content": system_prompt},
                    {"role": "user", "content": f"请基于以下信息生成内容：\n\n{context}"}
                ],
                "temperature": 0.7,
                "max_tokens": 4000,
                "stream": False
            }
            response = requests.post(url, headers=headers, json=payload, timeout=120)
            response.raise_for_status()
            resp_data = response.json()
            generated_content = resp_data["choices"][0]["message"]["content"]
            
        elif llm_provider == "ANTHROPIC":
            from anthropic import Anthropic
            api_key = os.getenv("ANTHROPIC_API_KEY")
            base_url = os.getenv("ANTHROPIC_BASE_URL")
            
            if not api_key:
                return jsonify({"success": False, "error": "ANTHROPIC_API_KEY 未配置"}), 500
            
            if base_url:
                client = Anthropic(api_key=api_key, base_url=base_url)
            else:
                client = Anthropic(api_key=api_key)
            
            model = os.getenv("ANTHROPIC_MODEL", "claude-3-5-sonnet-20241022")
            response = client.messages.create(
                model=model,
                max_tokens=4000,
                system=system_prompt,
                messages=[{"role": "user", "content": f"请基于以下信息生成内容：\n\n{context}"}]
            )
            generated_content = response.content[0].text if response.content else ""
            
        else:
            return jsonify({"success": False, "error": f"不支持的 LLM 提供商: {llm_provider}"}), 500
        
        # 保存生成的内容
        conn = sqlite3.connect(db.db_path)
        cursor = conn.cursor()
        
        # 获取当前版本号
        cursor.execute('''
            SELECT current_version FROM package_contents
            WHERE project_id = ? AND package_type = ?
        ''', (project_id, package_type))
        
        row = cursor.fetchone()
        new_version = (row[0] + 1) if row else 1
        
        # 保存版本历史
        word_count = len(generated_content.split()) if generated_content else 0
        cursor.execute('''
            INSERT INTO document_versions 
            (project_id, package_type, version, content, edit_type, edit_summary, editor, word_count, source_type)
            VALUES (?, ?, ?, ?, 'ai_generate', 'AI生成内容', 'ai', ?, 'ai_generated')
        ''', (project_id, package_type, new_version, generated_content, word_count))
        
        # 更新当前内容
        cursor.execute('''
            INSERT INTO package_contents 
            (project_id, package_type, current_version, content, status, last_edited_by, ai_generated, updated_at)
            VALUES (?, ?, ?, ?, 'draft', 'ai', 1, CURRENT_TIMESTAMP)
            ON CONFLICT(project_id, package_type) 
            DO UPDATE SET 
                current_version = excluded.current_version,
                content = excluded.content,
                status = 'draft',
                last_edited_by = 'ai',
                ai_generated = 1,
                updated_at = CURRENT_TIMESTAMP
        ''', (project_id, package_type, new_version, generated_content))
        
        conn.commit()
        conn.close()
        
        logger.info(f"AI生成材料包内容: {project_id}/{package_type} v{new_version}")
        
        return jsonify({
            "success": True,
            "content": generated_content,
            "version": new_version,
            "word_count": word_count
        })
        
    except Exception as e:
        logger.error(f"生成材料包内容失败: {e}")
        import traceback
        traceback.print_exc()
        return jsonify({"success": False, "error": str(e)}), 500


@copywriting_bp.route('/projects/<project_id>/packages/<package_type>/versions', methods=['GET'])
def get_package_versions(project_id, package_type):
    """获取材料包版本历史"""
    import sqlite3
    logger = _get_logger()
    
    try:
        _init_services()
        db = _services.get('db')
        
        if not db:
            return jsonify({"success": False, "error": "服务未初始化"}), 500
        
        conn = sqlite3.connect(db.db_path)
        conn.row_factory = sqlite3.Row
        cursor = conn.cursor()
        
        cursor.execute('''
            SELECT id, version, edit_type, edit_summary, editor, word_count, created_at
            FROM document_versions
            WHERE project_id = ? AND package_type = ?
            ORDER BY version DESC
        ''', (project_id, package_type))
        
        versions = []
        for row in cursor.fetchall():
            versions.append({
                "id": row['id'],
                "version": row['version'],
                "edit_type": row['edit_type'],
                "edit_summary": row['edit_summary'],
                "editor": row['editor'],
                "word_count": row['word_count'],
                "created_at": row['created_at']
            })
        
        conn.close()
        
        return jsonify({
            "success": True,
            "data": versions
        })
        
    except Exception as e:
        logger.error(f"获取版本历史失败: {e}")
        return jsonify({"success": False, "error": str(e)}), 500


@copywriting_bp.route('/projects/<project_id>/packages/<package_type>/versions/<int:version>', methods=['GET'])
def get_package_version_content(project_id, package_type, version):
    """获取特定版本的内容"""
    import sqlite3
    logger = _get_logger()
    
    try:
        _init_services()
        db = _services.get('db')
        
        if not db:
            return jsonify({"success": False, "error": "服务未初始化"}), 500
        
        conn = sqlite3.connect(db.db_path)
        conn.row_factory = sqlite3.Row
        cursor = conn.cursor()
        
        cursor.execute('''
            SELECT * FROM document_versions
            WHERE project_id = ? AND package_type = ? AND version = ?
        ''', (project_id, package_type, version))
        
        row = cursor.fetchone()
        conn.close()
        
        if row:
            return jsonify({
                "success": True,
                "data": {
                    "version": row['version'],
                    "content": row['content'],
                    "edit_type": row['edit_type'],
                    "edit_summary": row['edit_summary'],
                    "editor": row['editor'],
                    "word_count": row['word_count'],
                    "created_at": row['created_at']
                }
            })
        else:
            return jsonify({"success": False, "error": "版本不存在"}), 404
        
    except Exception as e:
        logger.error(f"获取版本内容失败: {e}")
        return jsonify({"success": False, "error": str(e)}), 500


@copywriting_bp.route('/projects/<project_id>/packages/<package_type>/rollback', methods=['POST'])
def rollback_package_version(project_id, package_type):
    """回滚到指定版本"""
    import sqlite3
    logger = _get_logger()
    
    try:
        _init_services()
        db = _services.get('db')
        
        if not db:
            return jsonify({"success": False, "error": "服务未初始化"}), 500
        
        data = request.get_json()
        target_version = data.get('version')
        
        if not target_version:
            return jsonify({"success": False, "error": "请指定目标版本"}), 400
        
        conn = sqlite3.connect(db.db_path)
        conn.row_factory = sqlite3.Row
        cursor = conn.cursor()
        
        # 获取目标版本内容
        cursor.execute('''
            SELECT content FROM document_versions
            WHERE project_id = ? AND package_type = ? AND version = ?
        ''', (project_id, package_type, target_version))
        
        row = cursor.fetchone()
        if not row:
            conn.close()
            return jsonify({"success": False, "error": "目标版本不存在"}), 404
        
        target_content = row['content']
        
        # 获取当前版本号
        cursor.execute('''
            SELECT current_version FROM package_contents
            WHERE project_id = ? AND package_type = ?
        ''', (project_id, package_type))
        
        current_row = cursor.fetchone()
        new_version = (current_row['current_version'] + 1) if current_row else 1
        
        # 保存新版本（回滚）
        word_count = len(target_content.split()) if target_content else 0
        cursor.execute('''
            INSERT INTO document_versions 
            (project_id, package_type, version, content, edit_type, edit_summary, editor, word_count)
            VALUES (?, ?, ?, ?, 'rollback', ?, 'user', ?)
        ''', (project_id, package_type, new_version, target_content, f"回滚到版本 {target_version}", word_count))
        
        # 更新当前内容
        cursor.execute('''
            UPDATE package_contents 
            SET current_version = ?, content = ?, last_edited_by = 'user', updated_at = CURRENT_TIMESTAMP
            WHERE project_id = ? AND package_type = ?
        ''', (new_version, target_content, project_id, package_type))
        
        conn.commit()
        conn.close()
        
        return jsonify({
            "success": True,
            "version": new_version,
            "message": f"已回滚到版本 {target_version}，保存为新版本 {new_version}"
        })
        
    except Exception as e:
        logger.error(f"回滚版本失败: {e}")
        return jsonify({"success": False, "error": str(e)}), 500


@copywriting_bp.route('/projects/<project_id>/packages/<package_type>/upload', methods=['POST'])
def upload_package_document(project_id, package_type):
    """上传文档并解析内容，创建新版本"""
    import sqlite3
    import tempfile
    import os
    logger = _get_logger()
    
    try:
        _init_services()
        db = _services.get('db')
        
        if not db:
            return jsonify({"success": False, "error": "服务未初始化"}), 500
        
        # 检查是否有文件
        if 'file' not in request.files:
            return jsonify({"success": False, "error": "请上传文件"}), 400
        
        file = request.files['file']
        if file.filename == '':
            return jsonify({"success": False, "error": "未选择文件"}), 400
        
        filename = file.filename
        file_ext = filename.rsplit('.', 1)[-1].lower() if '.' in filename else ''
        
        # 验证文件类型
        allowed_extensions = ['pdf', 'doc', 'docx', 'txt', 'md']
        if file_ext not in allowed_extensions:
            return jsonify({"success": False, "error": f"不支持的文件类型: {file_ext}"}), 400
        
        # 保存临时文件
        with tempfile.NamedTemporaryFile(delete=False, suffix=f'.{file_ext}') as tmp_file:
            file.save(tmp_file.name)
            tmp_path = tmp_file.name
        
        try:
            # 解析文件内容
            content = ""
            
            if file_ext in ['txt', 'md']:
                with open(tmp_path, 'r', encoding='utf-8') as f:
                    content = f.read()
            
            elif file_ext == 'docx':
                try:
                    from docx import Document
                    doc = Document(tmp_path)
                    paragraphs = [para.text for para in doc.paragraphs]
                    content = '\n\n'.join(paragraphs)
                except ImportError:
                    return jsonify({"success": False, "error": "服务器未安装 python-docx 库"}), 500
            
            elif file_ext == 'doc':
                # 尝试使用 antiword 或其他工具
                try:
                    import subprocess
                    result = subprocess.run(['antiword', tmp_path], capture_output=True, text=True)
                    if result.returncode == 0:
                        content = result.stdout
                    else:
                        return jsonify({"success": False, "error": "无法解析 .doc 文件，请转换为 .docx 格式"}), 400
                except FileNotFoundError:
                    return jsonify({"success": False, "error": "服务器无法解析 .doc 文件，请转换为 .docx 格式"}), 400
            
            elif file_ext == 'pdf':
                try:
                    import pdfplumber
                    with pdfplumber.open(tmp_path) as pdf:
                        pages_text = [page.extract_text() or '' for page in pdf.pages]
                        content = '\n\n'.join(pages_text)
                except ImportError:
                    try:
                        import PyPDF2
                        with open(tmp_path, 'rb') as f:
                            reader = PyPDF2.PdfReader(f)
                            pages_text = [page.extract_text() or '' for page in reader.pages]
                            content = '\n\n'.join(pages_text)
                    except ImportError:
                        return jsonify({"success": False, "error": "服务器未安装 PDF 解析库"}), 500
            
            if not content or not content.strip():
                return jsonify({"success": False, "error": "无法从文件中提取内容"}), 400
            
            # 计算字数
            word_count = len(content.split())
            
            # 获取当前版本号并创建新版本
            conn = sqlite3.connect(db.db_path)
            conn.row_factory = sqlite3.Row
            cursor = conn.cursor()
            
            cursor.execute('''
                SELECT current_version FROM package_contents
                WHERE project_id = ? AND package_type = ?
            ''', (project_id, package_type))
            
            current_row = cursor.fetchone()
            new_version = (current_row['current_version'] + 1) if current_row else 1
            
            # 插入版本记录
            cursor.execute('''
                INSERT INTO document_versions 
                (project_id, package_type, version, content, edit_type, edit_summary, editor, word_count, source_type, source_file)
                VALUES (?, ?, ?, ?, 'upload', ?, 'user', ?, 'uploaded', ?)
            ''', (project_id, package_type, new_version, content, f"上传文件: {filename}", word_count, filename))
            
            # 更新或插入当前内容
            if current_row:
                cursor.execute('''
                    UPDATE package_contents 
                    SET current_version = ?, content = ?, last_edited_by = 'user', 
                        ai_generated = 0, updated_at = CURRENT_TIMESTAMP
                    WHERE project_id = ? AND package_type = ?
                ''', (new_version, content, project_id, package_type))
            else:
                cursor.execute('''
                    INSERT INTO package_contents 
                    (project_id, package_type, current_version, content, status, last_edited_by, ai_generated)
                    VALUES (?, ?, ?, ?, 'draft', 'user', 0)
                ''', (project_id, package_type, new_version, content))
            
            conn.commit()
            conn.close()
            
            logger.info(f"文档上传成功: {project_id}/{package_type}, 版本 {new_version}")
            
            return jsonify({
                "success": True,
                "content": content,
                "version": new_version,
                "word_count": word_count,
                "source_file": filename
            })
            
        finally:
            # 清理临时文件
            if os.path.exists(tmp_path):
                os.remove(tmp_path)
        
    except Exception as e:
        logger.error(f"上传文档失败: {e}")
        return jsonify({"success": False, "error": str(e)}), 500


@copywriting_bp.route('/projects/<project_id>/packages/<package_type>/diff', methods=['GET'])
def get_package_diff(project_id, package_type):
    """获取两个版本之间的差异"""
    import sqlite3
    import difflib
    logger = _get_logger()
    
    try:
        _init_services()
        db = _services.get('db')
        
        if not db:
            return jsonify({"success": False, "error": "服务未初始化"}), 500
        
        v1 = request.args.get('v1', type=int)
        v2 = request.args.get('v2', type=int)
        
        if not v1 or not v2:
            return jsonify({"success": False, "error": "请指定 v1 和 v2 参数"}), 400
        
        conn = sqlite3.connect(db.db_path)
        conn.row_factory = sqlite3.Row
        cursor = conn.cursor()
        
        # 获取两个版本的内容
        cursor.execute('''
            SELECT version, content, edit_type, word_count, created_at, source_type, source_file
            FROM document_versions
            WHERE project_id = ? AND package_type = ? AND version IN (?, ?)
        ''', (project_id, package_type, v1, v2))
        
        rows = cursor.fetchall()
        conn.close()
        
        if len(rows) < 2:
            return jsonify({"success": False, "error": "找不到指定的版本"}), 404
        
        # 整理版本信息
        versions_data = {}
        for row in rows:
            versions_data[row['version']] = {
                'version': row['version'],
                'content': row['content'] or '',
                'edit_type': row['edit_type'],
                'word_count': row['word_count'],
                'created_at': row['created_at'],
                'source_type': row['source_type'] or row['edit_type'],
                'source_file': row['source_file']
            }
        
        v1_data = versions_data.get(v1, {'content': ''})
        v2_data = versions_data.get(v2, {'content': ''})
        
        # 生成行级差异
        v1_lines = v1_data['content'].split('\n')
        v2_lines = v2_data['content'].split('\n')
        
        diff_result = []
        matcher = difflib.SequenceMatcher(None, v1_lines, v2_lines)
        
        old_line_num = 0
        new_line_num = 0
        
        for tag, i1, i2, j1, j2 in matcher.get_opcodes():
            if tag == 'equal':
                for i in range(i2 - i1):
                    old_line_num += 1
                    new_line_num += 1
                    diff_result.append({
                        'type': 'equal',
                        'content': v1_lines[i1 + i],
                        'oldLineNum': old_line_num,
                        'newLineNum': new_line_num
                    })
            elif tag == 'delete':
                for i in range(i2 - i1):
                    old_line_num += 1
                    diff_result.append({
                        'type': 'delete',
                        'content': v1_lines[i1 + i],
                        'oldLineNum': old_line_num
                    })
            elif tag == 'insert':
                for j in range(j2 - j1):
                    new_line_num += 1
                    diff_result.append({
                        'type': 'insert',
                        'content': v2_lines[j1 + j],
                        'newLineNum': new_line_num
                    })
            elif tag == 'replace':
                for i in range(i2 - i1):
                    old_line_num += 1
                    diff_result.append({
                        'type': 'delete',
                        'content': v1_lines[i1 + i],
                        'oldLineNum': old_line_num
                    })
                for j in range(j2 - j1):
                    new_line_num += 1
                    diff_result.append({
                        'type': 'insert',
                        'content': v2_lines[j1 + j],
                        'newLineNum': new_line_num
                    })
        
        return jsonify({
            "success": True,
            "diff": diff_result,
            "v1_info": {
                'version': v1,
                'word_count': v1_data.get('word_count', 0),
                'created_at': v1_data.get('created_at'),
                'source_type': v1_data.get('source_type', 'manual')
            },
            "v2_info": {
                'version': v2,
                'word_count': v2_data.get('word_count', 0),
                'created_at': v2_data.get('created_at'),
                'source_type': v2_data.get('source_type', 'manual')
            },
            "v1_content": v1_data['content'],
            "v2_content": v2_data['content']
        })
        
    except Exception as e:
        logger.error(f"获取版本差异失败: {e}")
        return jsonify({"success": False, "error": str(e)}), 500


@copywriting_bp.route('/projects/<project_id>/packages/<package_type>/agent-config', methods=['GET'])
def get_agent_config(project_id, package_type):
    """获取Agent配置"""
    import sqlite3
    logger = _get_logger()
    
    try:
        _init_services()
        db = _services.get('db')
        
        if not db:
            return jsonify({"success": False, "error": "服务未初始化"}), 500
        
        conn = sqlite3.connect(db.db_path)
        conn.row_factory = sqlite3.Row
        cursor = conn.cursor()
        
        # 确保表存在
        cursor.execute('''
            CREATE TABLE IF NOT EXISTS agent_configs (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                project_id TEXT NOT NULL,
                package_type TEXT NOT NULL,
                system_prompt TEXT,
                user_prompt_template TEXT,
                custom_instructions TEXT,
                reference_doc_id TEXT,
                created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
                updated_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
                UNIQUE(project_id, package_type)
            )
        ''')
        conn.commit()
        
        cursor.execute('''
            SELECT * FROM agent_configs
            WHERE project_id = ? AND package_type = ?
        ''', (project_id, package_type))
        
        row = cursor.fetchone()
        conn.close()
        
        if row:
            return jsonify({
                "success": True,
                "data": {
                    "system_prompt": row['system_prompt'],
                    "user_prompt_template": row['user_prompt_template'],
                    "custom_instructions": row['custom_instructions'],
                    "reference_doc_id": row['reference_doc_id']
                }
            })
        else:
            return jsonify({
                "success": True,
                "data": {
                    "system_prompt": "",
                    "user_prompt_template": "",
                    "custom_instructions": "",
                    "reference_doc_id": None
                }
            })
        
    except Exception as e:
        logger.error(f"获取Agent配置失败: {e}")
        return jsonify({"success": False, "error": str(e)}), 500


@copywriting_bp.route('/projects/<project_id>/packages/<package_type>/agent-config', methods=['PUT'])
def save_agent_config(project_id, package_type):
    """保存Agent配置"""
    import sqlite3
    logger = _get_logger()
    
    try:
        _init_services()
        db = _services.get('db')
        
        if not db:
            return jsonify({"success": False, "error": "服务未初始化"}), 500
        
        data = request.get_json()
        
        conn = sqlite3.connect(db.db_path)
        cursor = conn.cursor()
        
        # 确保表存在
        cursor.execute('''
            CREATE TABLE IF NOT EXISTS agent_configs (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                project_id TEXT NOT NULL,
                package_type TEXT NOT NULL,
                system_prompt TEXT,
                user_prompt_template TEXT,
                custom_instructions TEXT,
                reference_doc_id TEXT,
                created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
                updated_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
                UNIQUE(project_id, package_type)
            )
        ''')
        
        cursor.execute('''
            INSERT INTO agent_configs 
            (project_id, package_type, system_prompt, user_prompt_template, custom_instructions, reference_doc_id, updated_at)
            VALUES (?, ?, ?, ?, ?, ?, CURRENT_TIMESTAMP)
            ON CONFLICT(project_id, package_type) 
            DO UPDATE SET 
                system_prompt = excluded.system_prompt,
                user_prompt_template = excluded.user_prompt_template,
                custom_instructions = excluded.custom_instructions,
                reference_doc_id = excluded.reference_doc_id,
                updated_at = CURRENT_TIMESTAMP
        ''', (
            project_id, 
            package_type, 
            data.get('system_prompt', ''),
            data.get('user_prompt_template', ''),
            data.get('custom_instructions', ''),
            data.get('reference_doc_id')
        ))
        
        conn.commit()
        conn.close()
        
        return jsonify({
            "success": True,
            "message": "Agent配置已保存"
        })
        
    except Exception as e:
        logger.error(f"保存Agent配置失败: {e}")
        return jsonify({"success": False, "error": str(e)}), 500


# ==================== 用户认证 API ====================

@copywriting_bp.route('/auth/register', methods=['POST'])
def auth_register():
    """用户注册"""
    logger = _get_logger()
    
    try:
        _init_services()
        db = _services.get('db')
        
        if not db:
            return jsonify({"success": False, "error": "服务未初始化"}), 500
        
        data = request.get_json()
        email = data.get('email')
        password_hash = data.get('password_hash')
        full_name = data.get('full_name')
        phone = data.get('phone')
        role = data.get('role', 'guest')
        
        if not email or not password_hash:
            return jsonify({"success": False, "error": "邮箱和密码为必填项"}), 400
        
        result = db.create_user(
            email=email,
            password_hash=password_hash,
            full_name=full_name,
            phone=phone,
            role=role
        )
        
        if result.get('success'):
            return jsonify(result)
        else:
            return jsonify(result), 400
        
    except Exception as e:
        logger.error(f"用户注册失败: {e}")
        return jsonify({"success": False, "error": str(e)}), 500


@copywriting_bp.route('/auth/user-by-email', methods=['POST'])
def auth_get_user_by_email():
    """根据邮箱获取用户"""
    logger = _get_logger()
    
    try:
        _init_services()
        db = _services.get('db')
        
        if not db:
            return jsonify({"success": False, "error": "服务未初始化"}), 500
        
        data = request.get_json()
        email = data.get('email')
        
        if not email:
            return jsonify({"success": False, "error": "邮箱为必填项"}), 400
        
        result = db.get_user_by_email(email)
        return jsonify(result)
        
    except Exception as e:
        logger.error(f"获取用户失败: {e}")
        return jsonify({"success": False, "error": str(e)}), 500


@copywriting_bp.route('/auth/user-by-id', methods=['POST'])
def auth_get_user_by_id():
    """根据ID获取用户"""
    logger = _get_logger()
    
    try:
        _init_services()
        db = _services.get('db')
        
        if not db:
            return jsonify({"success": False, "error": "服务未初始化"}), 500
        
        data = request.get_json()
        user_id = data.get('user_id')
        
        if not user_id:
            return jsonify({"success": False, "error": "用户ID为必填项"}), 400
        
        result = db.get_user_by_id(user_id)
        return jsonify(result)
        
    except Exception as e:
        logger.error(f"获取用户失败: {e}")
        return jsonify({"success": False, "error": str(e)}), 500


@copywriting_bp.route('/auth/session', methods=['POST'])
def auth_create_session():
    """创建用户会话"""
    logger = _get_logger()
    
    try:
        _init_services()
        db = _services.get('db')
        
        if not db:
            return jsonify({"success": False, "error": "服务未初始化"}), 500
        
        data = request.get_json()
        user_id = data.get('user_id')
        token = data.get('token')
        expires_at = data.get('expires_at')
        
        if not user_id or not token or not expires_at:
            return jsonify({"success": False, "error": "缺少必要参数"}), 400
        
        # 解析过期时间
        from datetime import datetime
        if isinstance(expires_at, str):
            expires_at = datetime.fromisoformat(expires_at.replace('Z', '+00:00'))
        
        result = db.create_session(user_id, token, expires_at)
        return jsonify(result)
        
    except Exception as e:
        logger.error(f"创建会话失败: {e}")
        return jsonify({"success": False, "error": str(e)}), 500


@copywriting_bp.route('/auth/session', methods=['DELETE'])
def auth_delete_session():
    """删除用户会话"""
    logger = _get_logger()
    
    try:
        _init_services()
        db = _services.get('db')
        
        if not db:
            return jsonify({"success": False, "error": "服务未初始化"}), 500
        
        data = request.get_json()
        token = data.get('token')
        
        if not token:
            return jsonify({"success": False, "error": "Token为必填项"}), 400
        
        result = db.delete_session(token)
        return jsonify(result)
        
    except Exception as e:
        logger.error(f"删除会话失败: {e}")
        return jsonify({"success": False, "error": str(e)}), 500


@copywriting_bp.route('/auth/validate-session', methods=['POST'])
def auth_validate_session():
    """验证会话"""
    logger = _get_logger()
    
    try:
        _init_services()
        db = _services.get('db')
        
        if not db:
            return jsonify({"success": False, "error": "服务未初始化"}), 500
        
        data = request.get_json()
        token = data.get('token')
        
        if not token:
            return jsonify({"success": False, "valid": False, "error": "Token为必填项"}), 400
        
        result = db.validate_session(token)
        return jsonify(result)
        
    except Exception as e:
        logger.error(f"验证会话失败: {e}")
        return jsonify({"success": False, "valid": False, "error": str(e)}), 500


@copywriting_bp.route('/auth/update-last-sign-in', methods=['POST'])
def auth_update_last_sign_in():
    """更新最后登录时间"""
    logger = _get_logger()
    
    try:
        _init_services()
        db = _services.get('db')
        
        if not db:
            return jsonify({"success": False, "error": "服务未初始化"}), 500
        
        data = request.get_json()
        user_id = data.get('user_id')
        
        if not user_id:
            return jsonify({"success": False, "error": "用户ID为必填项"}), 400
        
        result = db.update_last_sign_in(user_id)
        return jsonify(result)
        
    except Exception as e:
        logger.error(f"更新登录时间失败: {e}")
        return jsonify({"success": False, "error": str(e)}), 500


@copywriting_bp.route('/auth/users', methods=['GET'])
def auth_list_users():
    """列出所有用户"""
    logger = _get_logger()
    
    try:
        _init_services()
        db = _services.get('db')
        
        if not db:
            return jsonify({"success": False, "error": "服务未初始化"}), 500
        
        page = request.args.get('page', 1, type=int)
        page_size = request.args.get('page_size', 20, type=int)
        role = request.args.get('role')
        
        result = db.list_users(page=page, page_size=page_size, role=role)
        return jsonify(result)
        
    except Exception as e:
        logger.error(f"列出用户失败: {e}")
        return jsonify({"success": False, "error": str(e)}), 500


@copywriting_bp.route('/auth/users/<user_id>', methods=['PATCH'])
def auth_update_user(user_id):
    """更新用户信息"""
    logger = _get_logger()
    
    try:
        _init_services()
        db = _services.get('db')
        
        if not db:
            return jsonify({"success": False, "error": "服务未初始化"}), 500
        
        data = request.get_json()
        
        result = db.update_user(user_id, **data)
        return jsonify(result)
        
    except Exception as e:
        logger.error(f"更新用户失败: {e}")
        return jsonify({"success": False, "error": str(e)}), 500


@copywriting_bp.route('/auth/users/<user_id>', methods=['DELETE'])
def auth_delete_user(user_id):
    """删除用户"""
    logger = _get_logger()
    
    try:
        _init_services()
        db = _services.get('db')
        
        if not db:
            return jsonify({"success": False, "error": "服务未初始化"}), 500
        
        result = db.delete_user(user_id)
        return jsonify(result)
        
    except Exception as e:
        logger.error(f"删除用户失败: {e}")
        return jsonify({"success": False, "error": str(e)}), 500


# ==================== 评估记录 API ====================

@copywriting_bp.route('/assessments', methods=['POST'])
def save_assessment():
    """保存评估记录"""
    logger = _get_logger()
    
    try:
        _init_services()
        db = _services.get('db')
        
        if not db:
            return jsonify({"success": False, "error": "服务未初始化"}), 500
        
        data = request.get_json()
        
        result = db.save_assessment(data)
        return jsonify(result)
        
    except Exception as e:
        logger.error(f"保存评估记录失败: {e}")
        return jsonify({"success": False, "error": str(e)}), 500


@copywriting_bp.route('/assessments/<assessment_id>', methods=['GET'])
def get_assessment(assessment_id):
    """获取评估记录"""
    logger = _get_logger()
    
    try:
        _init_services()
        db = _services.get('db')
        
        if not db:
            return jsonify({"success": False, "error": "服务未初始化"}), 500
        
        result = db.get_assessment(assessment_id)
        return jsonify(result)
        
    except Exception as e:
        logger.error(f"获取评估记录失败: {e}")
        return jsonify({"success": False, "error": str(e)}), 500


@copywriting_bp.route('/assessments', methods=['GET'])
def list_assessments():
    """列出评估记录"""
    logger = _get_logger()
    
    try:
        _init_services()
        db = _services.get('db')
        
        if not db:
            return jsonify({"success": False, "error": "服务未初始化"}), 500
        
        user_id = request.args.get('user_id')
        page = request.args.get('page', 1, type=int)
        page_size = request.args.get('page_size', 20, type=int)
        
        result = db.list_assessments(user_id=user_id, page=page, page_size=page_size)
        return jsonify(result)
        
    except Exception as e:
        logger.error(f"列出评估记录失败: {e}")
        return jsonify({"success": False, "error": str(e)}), 500


# ==================== AI 文案助手 API ====================

@copywriting_bp.route('/assistant/skills', methods=['GET'])
def list_assistant_skills():
    """获取可用的 skills 列表"""
    logger = _get_logger()
    
    try:
        try:
            from ace_gtv.services.claude_code_service import get_claude_code_service
        except ImportError:
            from services.claude_code_service import get_claude_code_service
        
        service = get_claude_code_service()
        skills = service.list_skills()
        
        return jsonify({
            "success": True,
            "data": skills
        })
        
    except Exception as e:
        logger.error(f"获取 skills 列表失败: {e}")
        return jsonify({"success": False, "error": str(e)}), 500


@copywriting_bp.route('/assistant/skills/<skill_name>/content', methods=['GET'])
def get_skill_content(skill_name):
    """获取指定 skill 的内容"""
    logger = _get_logger()
    
    try:
        try:
            from ace_gtv.services.claude_code_service import get_claude_code_service
        except ImportError:
            from services.claude_code_service import get_claude_code_service
        
        service = get_claude_code_service()
        skill_info = service.available_skills.get(skill_name)
        
        if not skill_info:
            return jsonify({"success": False, "error": f"Skill not found: {skill_name}"}), 404
        
        # 读取 skill 内容
        from pathlib import Path
        skill_path = Path(skill_info.path)
        if skill_path.exists():
            content = skill_path.read_text(encoding='utf-8')
            return jsonify({
                "success": True,
                "data": {
                    "name": skill_info.name,
                    "display_name": skill_info.display_name,
                    "description": skill_info.description,
                    "content": content
                }
            })
        else:
            return jsonify({"success": False, "error": "Skill file not found"}), 404
        
    except Exception as e:
        logger.error(f"获取 skill 内容失败: {e}")
        return jsonify({"success": False, "error": str(e)}), 500


@copywriting_bp.route('/assistant/commands', methods=['GET'])
def list_assistant_commands():
    """获取可用的斜杠命令列表"""
    logger = _get_logger()
    
    try:
        try:
            from ace_gtv.services.command_scanner import get_command_scanner
        except ImportError:
            from services.command_scanner import get_command_scanner
        
        scanner = get_command_scanner()
        project_path = request.args.get('project_path')
        commands = scanner.list_commands(project_path)
        
        # 转换为可序列化的格式
        commands_data = []
        for cmd in commands:
            commands_data.append({
                "id": cmd.id,
                "name": cmd.name,
                "description": cmd.description,
                "category": cmd.category,
                "source": cmd.source,
                "argument_hint": cmd.argument_hint
            })
        
        return jsonify({
            "success": True,
            "data": commands_data
        })
        
    except Exception as e:
        logger.error(f"获取命令列表失败: {e}")
        return jsonify({"success": False, "error": str(e)}), 500


@copywriting_bp.route('/assistant/commands/<command_name>/content', methods=['GET'])
def get_command_content(command_name):
    """获取指定命令的内容"""
    logger = _get_logger()
    
    try:
        try:
            from ace_gtv.services.command_scanner import get_command_scanner
        except ImportError:
            from services.command_scanner import get_command_scanner
        
        scanner = get_command_scanner()
        project_path = request.args.get('project_path')
        command = scanner.get_command(command_name, project_path)
        
        if not command:
            return jsonify({"success": False, "error": f"Command not found: {command_name}"}), 404
        
        content = scanner.get_command_content(command)
        
        return jsonify({
            "success": True,
            "data": {
                "id": command.id,
                "name": command.name,
                "description": command.description,
                "category": command.category,
                "source": command.source,
                "argument_hint": command.argument_hint,
                "content": content
            }
        })
        
    except Exception as e:
        logger.error(f"获取命令内容失败: {e}")
        return jsonify({"success": False, "error": str(e)}), 500


@copywriting_bp.route('/assistant/memory', methods=['GET'])
def get_memory_info():
    """获取项目记忆文件信息"""
    logger = _get_logger()
    
    try:
        try:
            from ace_gtv.services.memory_loader import get_memory_loader
        except ImportError:
            from services.memory_loader import get_memory_loader
        
        loader = get_memory_loader()
        project_path = request.args.get('project_path')
        memory_info = loader.get_memory_files(project_path)
        
        return jsonify({
            "success": True,
            "data": memory_info
        })
        
    except Exception as e:
        logger.error(f"获取记忆信息失败: {e}")
        return jsonify({"success": False, "error": str(e)}), 500


@copywriting_bp.route('/assistant/status', methods=['GET'])
def general_assistant_status():
    """
    获取 AI 助手服务状态（通用端点，不需要 project_id）
    
    返回：
    - current_mode: 当前执行模式 (ask/agent/plan)
    - available_modes: 可用模式列表
    - cli_available: CLI 是否可用
    - cli_path: CLI 路径
    - skills_count: 可用技能数量
    - llm_provider: API 模式使用的 LLM 提供商
    """
    try:
        try:
            from ace_gtv.services.claude_code_service import get_claude_code_service
        except ImportError:
            from services.claude_code_service import get_claude_code_service
        
        service = get_claude_code_service()
        status = service.get_status()
        
        return jsonify({
            "success": True,
            "data": status
        })
        
    except Exception as e:
        return jsonify({"success": False, "error": str(e)}), 500


@copywriting_bp.route('/assistant/chat', methods=['POST'])
def general_assistant_chat():
    """
    AI 助手对话 API（通用端点，不需要 project_id）
    
    支持流式响应（SSE）
    """
    logger = _get_logger()
    
    try:
        data = request.get_json()
        message = data.get('message', '')
        skill = data.get('skill')
        stream = data.get('stream', True)
        mode = data.get('mode', 'ask')
        project_context = data.get('project_context', {})
        
        if not message:
            return jsonify({"success": False, "error": "消息不能为空"}), 400
        
        # 构建上下文
        context = {
            **project_context,
            "mode": mode
        }
        
        # 如果请求流式响应
        if stream:
            from flask import Response
            
            def generate():
                try:
                    from ace_gtv.services.claude_code_service import get_claude_code_service
                except ImportError:
                    from services.claude_code_service import get_claude_code_service
                
                service = get_claude_code_service()
                # 使用请求中的 mode 参数
                current_mode = mode
                
                # 发送开始信号
                yield f"data: {json.dumps({'type': 'start', 'mode': current_mode, 'skill': skill})}\n\n"
                
                for chunk in service.execute_with_skill(
                    prompt=message,
                    skill_name=skill,
                    context=context,
                    stream=True,
                    mode=current_mode  # 传递 mode 参数
                ):
                    if chunk.startswith('['):
                        yield f"data: {json.dumps({'type': 'log', 'content': chunk})}\n\n"
                    else:
                        yield f"data: {json.dumps({'type': 'content', 'content': chunk})}\n\n"
                
                yield f"data: {json.dumps({'type': 'done', 'skill': skill, 'mode': current_mode})}\n\n"
            
            return Response(
                generate(),
                mimetype='text/event-stream',
                headers={
                    'Cache-Control': 'no-cache',
                    'Connection': 'keep-alive',
                    'X-Accel-Buffering': 'no'
                }
            )
        
        # 非流式响应
        try:
            from ace_gtv.services.claude_code_service import get_claude_code_service
        except ImportError:
            from services.claude_code_service import get_claude_code_service
        
        service = get_claude_code_service()
        logger.info(f"[助手对话] 非流式响应，mode={mode}, message长度={len(message)}")
        
        result_chunks = []
        for chunk in service.execute_with_skill(
            prompt=message,
            skill_name=skill,
            context=context,
            stream=False,
            mode=mode  # 传递 mode 参数
        ):
            result_chunks.append(chunk)
            logger.debug(f"[助手对话] 接收 chunk: {len(chunk)} 字符")
        
        logger.info(f"[助手对话] 完成，共 {len(result_chunks)} 个 chunk，总长度 {sum(len(c) for c in result_chunks)}")
        return jsonify({
            "success": True,
            "data": {
                "content": "".join(result_chunks),
                "skill": skill,
                "mode": mode
            }
        })
        
    except Exception as e:
        logger.error(f"助手对话失败: {e}")
        import traceback
        traceback.print_exc()
        return jsonify({"success": False, "error": str(e)}), 500


@copywriting_bp.route('/projects/<project_id>/assistant/status', methods=['GET'])
def assistant_status(project_id):
    """
    获取 AI 助手服务状态
    
    返回：
    - execution_mode: 当前执行模式 (cli/api/auto)
    - cli_available: CLI 是否可用
    - cli_path: CLI 路径
    - skills_count: 可用技能数量
    - llm_provider: API 模式使用的 LLM 提供商
    """
    try:
        try:
            from ace_gtv.services.claude_code_service import get_claude_code_service
        except ImportError:
            from services.claude_code_service import get_claude_code_service
        
        service = get_claude_code_service()
        status = service.get_status()
        
        return jsonify({
            "success": True,
            "data": status
        })
        
    except Exception as e:
        return jsonify({"success": False, "error": str(e)}), 500


@copywriting_bp.route('/projects/<project_id>/assistant/mode', methods=['POST'])
def set_assistant_mode(project_id):
    """
    设置 AI 助手执行模式
    
    请求体:
    {
        "mode": "ask" | "agent" | "plan"
    }
    
    模式说明:
    - ask: 快速问答模式，直接调用 AI API
    - agent: 智能代理模式，使用 Claude CLI，支持复杂任务
    - plan: 规划模式，分析需求并生成详细计划
    
    注意：此设置会修改环境变量，影响当前会话的所有后续请求
    """
    import os
    
    try:
        data = request.get_json()
        mode = data.get('mode', '').lower()
        
        valid_modes = ['ask', 'agent', 'plan', 'auto']
        if mode not in valid_modes:
            return jsonify({
                "success": False,
                "error": f"无效的模式: {mode}，可选值: {', '.join(valid_modes)}"
            }), 400
        
        # 检查 agent 模式是否可用
        try:
            from ace_gtv.services.claude_code_service import get_claude_code_service
        except ImportError:
            from services.claude_code_service import get_claude_code_service
        
        service = get_claude_code_service()
        
        if mode == 'agent' and not service.is_cli_available():
            return jsonify({
                "success": False,
                "error": "Agent 模式不可用：Claude CLI 未安装"
            }), 400
        
        # 设置环境变量
        os.environ['AI_EXECUTION_MODE'] = mode
        
        status = service.get_status()
        
        # 模式名称映射
        mode_names = {
            'ask': 'Ask（快速问答）',
            'agent': 'Agent（智能代理）',
            'plan': 'Plan（规划模式）',
            'auto': 'Auto（自动选择）'
        }
        
        return jsonify({
            "success": True,
            "message": f"已切换到 {mode_names.get(mode, mode)} 模式",
            "data": status
        })
        
    except Exception as e:
        return jsonify({"success": False, "error": str(e)}), 500


@copywriting_bp.route('/projects/<project_id>/assistant/chat', methods=['POST'])
def assistant_chat(project_id):
    """
    AI 文案助手对话 API
    
    支持流式响应（SSE）
    """
    logger = _get_logger()
    
    try:
        _init_services()
        db = _services.get('db')
        
        if not db:
            return jsonify({"success": False, "error": "服务未初始化"}), 500
        
        data = request.get_json()
        message = data.get('message', '')
        skill = data.get('skill')  # 可选，手动指定的 skill
        mode = data.get('mode', 'ask')  # 执行模式: ask/agent/plan
        document_context = data.get('document_context')
        active_document = data.get('active_document')
        project_context = data.get('project_context', {})
        conversation_history = data.get('conversation_history', [])
        stream = data.get('stream', True)
        
        if not message:
            return jsonify({"success": False, "error": "消息不能为空"}), 400
        
        # 获取项目信息
        project_result = db.get_project(project_id)
        client_name = ""
        if project_result.get('success') and project_result.get('data'):
            client_name = project_result['data'].get('client_name', '')
        
        # 自动检测 skill（如果未指定）
        if not skill:
            try:
                from ace_gtv.services.skill_router import auto_detect_skill
            except ImportError:
                from services.skill_router import auto_detect_skill
            skill = auto_detect_skill(message, {
                "active_document": active_document,
                "client_name": client_name
            })
            logger.info(f"自动检测 skill: {skill}")
        
        # 构建上下文
        context = {
            "project_id": project_id,
            "client_name": client_name,
            "active_document": active_document,
            "document_context": document_context,
            "project_context": project_context,
            "conversation_history": conversation_history
        }
        
        # 如果请求流式响应
        if stream:
            from flask import Response
            
            # 使用请求中的 mode 参数
            exec_mode = mode  # 从请求参数获取
            
            def generate():
                try:
                    from ace_gtv.services.claude_code_service import get_claude_code_service
                except ImportError:
                    from services.claude_code_service import get_claude_code_service
                
                service = get_claude_code_service()
                
                # 发送开始信号，包含执行模式
                yield f"data: {json.dumps({'type': 'start', 'mode': exec_mode, 'skill': skill})}\n\n"
                
                # 如果是 CLI 模式，发送日志信息
                if exec_mode in ['agent', 'cli']:
                    cli_path_log = f'[系统] 使用 Agent 模式执行，Claude CLI 路径: {service._cli_path}'
                    yield f"data: {json.dumps({'type': 'log', 'content': cli_path_log})}\n\n"
                    skill_name = skill or '自动检测'
                    skill_log = f'[系统] 技能: {skill_name}'
                    yield f"data: {json.dumps({'type': 'log', 'content': skill_log})}\n\n"
                    yield f"data: {json.dumps({'type': 'log', 'content': '[系统] 正在执行...'})}\n\n"
                
                for chunk in service.execute_with_skill(
                    prompt=message,
                    skill_name=skill,
                    context=context,
                    stream=True,
                    mode=exec_mode  # 传递 mode 参数
                ):
                    # 判断是日志还是内容
                    if chunk.startswith('['):
                        # 这是日志/状态信息
                        yield f"data: {json.dumps({'type': 'log', 'content': chunk})}\n\n"
                    else:
                        # 这是实际内容
                        yield f"data: {json.dumps({'type': 'content', 'content': chunk})}\n\n"
                
                # 发送完成信号
                yield f"data: {json.dumps({'type': 'done', 'skill': skill, 'mode': exec_mode})}\n\n"
            
            return Response(
                generate(),
                mimetype='text/event-stream',
                headers={
                    'Cache-Control': 'no-cache',
                    'Connection': 'keep-alive',
                    'X-Accel-Buffering': 'no'
                }
            )
        
        # 非流式响应
        try:
            from ace_gtv.services.claude_code_service import get_claude_code_service
        except ImportError:
            from services.claude_code_service import get_claude_code_service
        
        service = get_claude_code_service()
        
        result_chunks = []
        for chunk in service.execute_with_skill(
            prompt=message,
            skill_name=skill,
            context=context,
            stream=False,
            mode=mode  # 传递 mode 参数
        ):
            result_chunks.append(chunk)
        
        return jsonify({
            "success": True,
            "message": "".join(result_chunks),
            "skill": skill,
            "mode": mode
        })
        
    except Exception as e:
        logger.error(f"助手对话失败: {e}")
        import traceback
        traceback.print_exc()
        return jsonify({"success": False, "error": str(e)}), 500


@copywriting_bp.route('/projects/<project_id>/assistant/suggest', methods=['POST'])
def assistant_suggest(project_id):
    """
    获取文档修改建议
    
    基于当前文档内容和用户需求，生成修改建议
    """
    logger = _get_logger()
    
    try:
        _init_services()
        db = _services.get('db')
        
        if not db:
            return jsonify({"success": False, "error": "服务未初始化"}), 500
        
        data = request.get_json()
        document_type = data.get('document_type')
        document_content = data.get('document_content', '')
        user_request = data.get('request', '')
        
        if not document_type:
            return jsonify({"success": False, "error": "文档类型不能为空"}), 400
        
        # 获取项目信息
        project_result = db.get_project(project_id)
        client_name = ""
        if project_result.get('success') and project_result.get('data'):
            client_name = project_result['data'].get('client_name', '')
        
        # 构建提示
        prompt = f"""请分析以下{document_type}文档，并提供具体的修改建议。

客户名称: {client_name}
文档类型: {document_type}
用户需求: {user_request or "请提供改进建议"}

当前文档内容:
{document_content[:3000]}

请以 JSON 格式返回修改建议，包含以下字段:
- suggestions: 建议列表，每个建议包含:
  - type: "edit" | "add" | "delete"
  - original_text: 原文（如适用）
  - suggested_text: 建议内容
  - reason: 修改理由
"""
        
        try:
            from ace_gtv.services.claude_code_service import get_claude_code_service
        except ImportError:
            from services.claude_code_service import get_claude_code_service
        
        service = get_claude_code_service()
        
        result_chunks = []
        for chunk in service.execute_with_skill(
            prompt=prompt,
            skill_name="recommendations-generation",
            context={"client_name": client_name, "active_document": document_type},
            stream=False
        ):
            result_chunks.append(chunk)
        
        result_text = "".join(result_chunks)
        
        # 尝试解析 JSON
        suggestions = []
        try:
            import re
            json_match = re.search(r'\{[\s\S]*\}', result_text)
            if json_match:
                parsed = json.loads(json_match.group())
                suggestions = parsed.get('suggestions', [])
        except json.JSONDecodeError:
            # 如果无法解析，返回原始文本作为单个建议
            suggestions = [{
                "type": "edit",
                "suggested_text": result_text,
                "reason": "AI 分析结果"
            }]
        
        return jsonify({
            "success": True,
            "data": {
                "suggestions": suggestions,
                "raw_response": result_text
            }
        })
        
    except Exception as e:
        logger.error(f"获取建议失败: {e}")
        return jsonify({"success": False, "error": str(e)}), 500


@copywriting_bp.route('/projects/<project_id>/assistant/apply-suggestion', methods=['POST'])
def assistant_apply_suggestion(project_id):
    """
    应用建议到文档
    
    将 AI 建议的修改应用到指定文档
    """
    logger = _get_logger()
    
    try:
        _init_services()
        db = _services.get('db')
        
        if not db:
            return jsonify({"success": False, "error": "服务未初始化"}), 500
        
        data = request.get_json()
        document_type = data.get('document_type')
        suggestion = data.get('suggestion', {})
        
        if not document_type or not suggestion:
            return jsonify({"success": False, "error": "参数不完整"}), 400
        
        # 获取当前文档内容
        package_result = db.get_package(project_id, document_type)
        if not package_result.get('success'):
            return jsonify({"success": False, "error": "获取文档失败"}), 404
        
        current_content = package_result.get('data', {}).get('content', '')
        
        # 应用建议
        new_content = current_content
        suggestion_type = suggestion.get('type', 'edit')
        original_text = suggestion.get('original_text', '')
        suggested_text = suggestion.get('suggested_text', '')
        
        if suggestion_type == 'edit' and original_text:
            new_content = current_content.replace(original_text, suggested_text)
        elif suggestion_type == 'add':
            new_content = current_content + "\n\n" + suggested_text
        elif suggestion_type == 'delete' and original_text:
            new_content = current_content.replace(original_text, '')
        
        # 保存更新后的文档
        update_result = db.update_package(project_id, document_type, {
            'content': new_content
        })
        
        if not update_result.get('success'):
            return jsonify({"success": False, "error": "保存失败"}), 500
        
        return jsonify({
            "success": True,
            "data": {
                "document_type": document_type,
                "new_content": new_content,
                "applied_suggestion": suggestion
            }
        })
        
    except Exception as e:
        logger.error(f"应用建议失败: {e}")
        return jsonify({"success": False, "error": str(e)}), 500


# ==================== Cloud CLI 管理 API ====================

def _get_cloudcli_manager():
    """获取 CloudCLI 管理器实例"""
    if 'cloudcli_manager' not in _services:
        try:
            from services.cloudcli_manager import get_cloudcli_manager
            _services['cloudcli_manager'] = get_cloudcli_manager()
        except ImportError:
            try:
                from ace_gtv.services.cloudcli_manager import get_cloudcli_manager
                _services['cloudcli_manager'] = get_cloudcli_manager()
            except ImportError:
                return None
    return _services['cloudcli_manager']


@copywriting_bp.route('/cloudcli/status', methods=['GET'])
def cloudcli_status():
    """
    获取 Cloud CLI 服务状态
    
    返回:
    {
        "success": true,
        "data": {
            "status": "running" | "stopped" | "starting" | "error" | "not_installed",
            "running": true | false,
            "url": "http://localhost:3001",
            "port": 3001,
            "host": "localhost",
            "auto_start": true,
            "error": null | "error message",
            "pid": 12345 | null
        }
    }
    """
    logger = _get_logger()
    
    try:
        manager = _get_cloudcli_manager()
        if manager is None:
            return jsonify({
                "success": False,
                "error": "Cloud CLI 管理器未初始化"
            }), 500
        
        status = manager.get_status()
        return jsonify({
            "success": True,
            "data": status
        })
        
    except Exception as e:
        logger.error(f"获取 Cloud CLI 状态失败: {e}")
        return jsonify({"success": False, "error": str(e)}), 500


@copywriting_bp.route('/cloudcli/start', methods=['POST'])
def cloudcli_start():
    """
    启动 Cloud CLI 服务
    
    请求体（可选）:
    {
        "force": false  // 是否强制重启
    }
    
    返回:
    {
        "success": true,
        "data": {
            "status": "running",
            "url": "http://localhost:3001",
            "message": "Cloud CLI 服务启动成功"
        }
    }
    """
    logger = _get_logger()
    
    try:
        manager = _get_cloudcli_manager()
        if manager is None:
            return jsonify({
                "success": False,
                "error": "Cloud CLI 管理器未初始化"
            }), 500
        
        data = request.get_json() or {}
        force = data.get('force', False)
        
        result = manager.start(force=force)
        
        if result.get('success'):
            return jsonify({
                "success": True,
                "data": result
            })
        else:
            return jsonify({
                "success": False,
                "error": result.get('error', '启动失败')
            }), 500
        
    except Exception as e:
        logger.error(f"启动 Cloud CLI 失败: {e}")
        return jsonify({"success": False, "error": str(e)}), 500


@copywriting_bp.route('/cloudcli/stop', methods=['POST'])
def cloudcli_stop():
    """
    停止 Cloud CLI 服务
    
    返回:
    {
        "success": true,
        "data": {
            "status": "stopped",
            "message": "Cloud CLI 服务已停止"
        }
    }
    """
    logger = _get_logger()
    
    try:
        manager = _get_cloudcli_manager()
        if manager is None:
            return jsonify({
                "success": False,
                "error": "Cloud CLI 管理器未初始化"
            }), 500
        
        result = manager.stop()
        
        return jsonify({
            "success": True,
            "data": result
        })
        
    except Exception as e:
        logger.error(f"停止 Cloud CLI 失败: {e}")
        return jsonify({"success": False, "error": str(e)}), 500


@copywriting_bp.route('/cloudcli/config', methods=['GET'])
def cloudcli_config():
    """
    获取 Cloud CLI 配置信息
    
    返回:
    {
        "success": true,
        "data": {
            "port": 3001,
            "host": "localhost",
            "auto_start": true,
            "startup_timeout": 30,
            "health_check_interval": 5,
            "cli_path": "/usr/local/bin/cloudcli",
            "use_npx": false,
            "installed": true
        }
    }
    """
    logger = _get_logger()
    
    try:
        manager = _get_cloudcli_manager()
        if manager is None:
            return jsonify({
                "success": False,
                "error": "Cloud CLI 管理器未初始化"
            }), 500
        
        config = manager.get_config()
        return jsonify({
            "success": True,
            "data": config
        })
        
    except Exception as e:
        logger.error(f"获取 Cloud CLI 配置失败: {e}")
        return jsonify({"success": False, "error": str(e)}), 500


@copywriting_bp.route('/cloudcli/ensure', methods=['POST'])
def cloudcli_ensure():
    """
    确保 Cloud CLI 服务正在运行
    如果配置了自动启动且服务未运行，将自动启动服务
    
    返回:
    {
        "success": true,
        "data": {
            "status": "running",
            "url": "http://localhost:3001"
        }
    }
    """
    logger = _get_logger()
    
    try:
        manager = _get_cloudcli_manager()
        if manager is None:
            return jsonify({
                "success": False,
                "error": "Cloud CLI 管理器未初始化"
            }), 500
        
        result = manager.ensure_running()
        
        return jsonify({
            "success": True,
            "data": result
        })
        
    except Exception as e:
        logger.error(f"确保 Cloud CLI 运行失败: {e}")
        return jsonify({"success": False, "error": str(e)}), 500


# ========== Claude Task API ==========

def _get_task_service():
    """获取 Claude Task 服务"""
    if 'task_service' not in _services:
        try:
            from ace_gtv.services.claude_task_service import get_task_service
        except ImportError:
            from services.claude_task_service import get_task_service
        _services['task_service'] = get_task_service()
    return _services['task_service']


def _get_workspace_service():
    """获取项目工作空间服务"""
    if 'workspace_service' not in _services:
        try:
            from ace_gtv.services.project_workspace_service import get_workspace_service
        except ImportError:
            from services.project_workspace_service import get_workspace_service
        _services['workspace_service'] = get_workspace_service()
    return _services['workspace_service']


# ========== Project Workspace API ==========

@copywriting_bp.route('/projects/<project_id>/workspace', methods=['GET'])
def get_project_workspace(project_id):
    """获取项目工作空间信息"""
    logger = _get_logger()
    
    try:
        service = _get_workspace_service()
        info = service.get_workspace_info(project_id)
        
        return jsonify({
            "success": True,
            "data": info
        })
        
    except Exception as e:
        logger.error(f"获取工作空间信息失败: {e}")
        return jsonify({"success": False, "error": str(e)}), 500


@copywriting_bp.route('/projects/<project_id>/workspace', methods=['POST'])
def prepare_project_workspace(project_id):
    """
    准备项目工作空间
    
    创建独立的工作目录，复制材料和技能文件
    """
    logger = _get_logger()
    
    try:
        data = request.get_json() or {}
        force = data.get('force', False)
        copy_uploads = data.get('copy_uploads', True)
        
        # 获取项目信息
        _init_services()
        db = _services.get('db')
        project_info = {}
        
        if db:
            project_result = db.get_project(project_id)
            if project_result.get('success') and project_result.get('data'):
                project_info = project_result['data']
        
        # 准备工作空间
        workspace_service = _get_workspace_service()
        
        if force:
            workspace_service.cleanup_workspace(project_id)
        
        result = workspace_service.prepare_workspace(
            project_id=project_id,
            project_info=project_info,
            copy_uploads=copy_uploads
        )
        
        return jsonify({
            "success": True,
            "data": result,
            "message": "工作空间已准备就绪"
        })
        
    except Exception as e:
        logger.error(f"准备工作空间失败: {e}")
        import traceback
        traceback.print_exc()
        return jsonify({"success": False, "error": str(e)}), 500


@copywriting_bp.route('/projects/<project_id>/workspace', methods=['DELETE'])
def cleanup_project_workspace(project_id):
    """清理项目工作空间"""
    logger = _get_logger()
    
    try:
        service = _get_workspace_service()
        result = service.cleanup_workspace(project_id)
        
        return jsonify({
            "success": result,
            "message": "工作空间已清理" if result else "工作空间不存在"
        })
        
    except Exception as e:
        logger.error(f"清理工作空间失败: {e}")
        return jsonify({"success": False, "error": str(e)}), 500


@copywriting_bp.route('/projects/<project_id>/workspace/materials', methods=['POST'])
def copy_materials_to_workspace(project_id):
    """复制材料到工作空间"""
    logger = _get_logger()
    
    try:
        data = request.get_json() or {}
        source_files = data.get('source_files', [])
        
        service = _get_workspace_service()
        
        # 确保工作空间存在
        if not service.workspace_exists(project_id):
            return jsonify({
                "success": False,
                "error": "工作空间不存在，请先创建工作空间"
            }), 400
        
        result = service.copy_materials(
            project_id=project_id,
            source_files=source_files,
            from_upload_dir=True
        )
        
        return jsonify({
            "success": True,
            "data": result
        })
        
    except Exception as e:
        logger.error(f"复制材料失败: {e}")
        return jsonify({"success": False, "error": str(e)}), 500


# ========== Claude Bridge API (旧版，保留兼容) ==========

def _get_bridge_service():
    """获取 Claude Bridge 服务"""
    if 'bridge_service' not in _services:
        try:
            from ace_gtv.services.claude_bridge_service import get_bridge_service
        except ImportError:
            from services.claude_bridge_service import get_bridge_service
        _services['bridge_service'] = get_bridge_service()
    return _services['bridge_service']


@copywriting_bp.route('/claude-bridge/sessions', methods=['POST'])
def create_claude_session():
    """
    创建 Claude Code 会话
    
    启动一个常驻的 Claude Code 进程，支持双向通信
    """
    logger = _get_logger()
    
    try:
        data = request.get_json() or {}
        session_id = data.get('session_id', f"session_{int(time.time() * 1000)}")
        cwd = data.get('cwd', os.getcwd())
        permission_mode = data.get('permission_mode', 'default')
        
        bridge = _get_bridge_service()
        session = bridge.create_session(
            session_id=session_id,
            cwd=cwd,
            permission_mode=permission_mode
        )
        
        return jsonify({
            "success": True,
            "data": bridge.get_session_status(session_id)
        })
        
    except Exception as e:
        logger.error(f"创建 Claude 会话失败: {e}")
        return jsonify({"success": False, "error": str(e)}), 500


@copywriting_bp.route('/claude-bridge/sessions', methods=['GET'])
def list_claude_sessions():
    """列出所有 Claude Code 会话"""
    logger = _get_logger()
    
    try:
        bridge = _get_bridge_service()
        sessions = bridge.list_sessions()
        
        return jsonify({
            "success": True,
            "data": sessions
        })
        
    except Exception as e:
        logger.error(f"获取会话列表失败: {e}")
        return jsonify({"success": False, "error": str(e)}), 500


@copywriting_bp.route('/claude-bridge/sessions/<session_id>', methods=['GET'])
def get_claude_session(session_id):
    """获取会话状态"""
    logger = _get_logger()
    
    try:
        bridge = _get_bridge_service()
        status = bridge.get_session_status(session_id)
        
        if not status:
            return jsonify({"success": False, "error": "会话不存在"}), 404
        
        return jsonify({
            "success": True,
            "data": status
        })
        
    except Exception as e:
        logger.error(f"获取会话状态失败: {e}")
        return jsonify({"success": False, "error": str(e)}), 500


@copywriting_bp.route('/claude-bridge/sessions/<session_id>', methods=['DELETE'])
def close_claude_session(session_id):
    """关闭会话"""
    logger = _get_logger()
    
    try:
        bridge = _get_bridge_service()
        result = bridge.close_session(session_id)
        
        return jsonify({
            "success": result,
            "message": "会话已关闭" if result else "会话不存在"
        })
        
    except Exception as e:
        logger.error(f"关闭会话失败: {e}")
        return jsonify({"success": False, "error": str(e)}), 500


@copywriting_bp.route('/claude-bridge/sessions/<session_id>/messages', methods=['POST'])
def send_claude_message(session_id):
    """
    向 Claude 会话发送消息
    
    发送消息到常驻的 Claude Code 进程
    支持发送空消息（用于回车键）
    """
    logger = _get_logger()
    
    try:
        data = request.get_json()
        message = data.get('message', '')
        # 允许空消息（用于发送回车）
        
        bridge = _get_bridge_service()
        result = bridge.send_message(session_id, message)
        
        if not result:
            return jsonify({"success": False, "error": "发送失败，会话可能已关闭"}), 400
        
        return jsonify({
            "success": True,
            "message": "消息已发送"
        })
        
    except Exception as e:
        logger.error(f"发送消息失败: {e}")
        return jsonify({"success": False, "error": str(e)}), 500


@copywriting_bp.route('/claude-bridge/sessions/<session_id>/stream', methods=['GET'])
def stream_claude_output(session_id):
    """
    流式获取 Claude 输出（SSE）
    
    实时推送 Claude Code 的输出到客户端
    """
    from flask import Response
    logger = _get_logger()
    
    def generate():
        bridge = _get_bridge_service()
        timeout = request.args.get('timeout', 300, type=int)
        
        for output in bridge.stream_output(session_id, timeout=timeout):
            yield f"data: {json.dumps(output, ensure_ascii=False)}\n\n"
    
    return Response(
        generate(),
        mimetype='text/event-stream',
        headers={
            'Cache-Control': 'no-cache',
            'Connection': 'keep-alive',
            'X-Accel-Buffering': 'no'
        }
    )


@copywriting_bp.route('/claude-bridge/sessions/<session_id>/output', methods=['GET'])
def get_claude_output(session_id):
    """
    获取 Claude 输出（轮询方式）
    
    非阻塞获取输出队列中的消息
    """
    logger = _get_logger()
    
    try:
        bridge = _get_bridge_service()
        timeout = request.args.get('timeout', 0.5, type=float)
        max_messages = request.args.get('max', 10, type=int)
        
        messages = []
        for _ in range(max_messages):
            output = bridge.get_output(session_id, timeout=timeout)
            if output:
                messages.append(output)
            else:
                break
        
        return jsonify({
            "success": True,
            "data": messages,
            "has_more": len(messages) == max_messages
        })
        
    except Exception as e:
        logger.error(f"获取输出失败: {e}")
        return jsonify({"success": False, "error": str(e)}), 500


@copywriting_bp.route('/claude-bridge/sessions/<session_id>/full-output', methods=['GET'])
def get_claude_full_output(session_id):
    """
    获取会话的完整输出
    """
    logger = _get_logger()
    
    try:
        bridge = _get_bridge_service()
        output = bridge.get_full_output(session_id)
        
        if output is None:
            return jsonify({"success": False, "error": "会话不存在"}), 404
        
        return jsonify({
            "success": True,
            "data": {
                "output": output,
                "length": len(output)
            }
        })
        
    except Exception as e:
        logger.error(f"获取完整输出失败: {e}")
        return jsonify({"success": False, "error": str(e)}), 500


@copywriting_bp.route('/claude-bridge/sessions/<session_id>/clear', methods=['POST'])
def clear_claude_output(session_id):
    """
    清空会话输出缓冲
    """
    logger = _get_logger()
    
    try:
        bridge = _get_bridge_service()
        result = bridge.clear_output_buffer(session_id)
        
        return jsonify({
            "success": result,
            "message": "输出已清空" if result else "会话不存在"
        })
        
    except Exception as e:
        logger.error(f"清空输出失败: {e}")
        return jsonify({"success": False, "error": str(e)}), 500


# 需要导入 time 模块
import time


# ========== Claude Task API (新版任务队列架构) ==========

@copywriting_bp.route('/claude-tasks', methods=['POST'])
def create_claude_task():
    """
    创建 Claude 任务
    
    请求体:
    {
        "prompt": "用户输入",
        "cwd": "工作目录（可选）",
        "async": true  // 是否异步执行
    }
    
    异步模式：立即返回任务ID，后台执行，通过 SSE 获取输出
    同步模式：等待执行完成后返回结果
    """
    logger = _get_logger()
    
    try:
        data = request.get_json()
        prompt = data.get('prompt', '')
        cwd = data.get('cwd', os.getcwd())
        async_mode = data.get('async', True)
        
        if not prompt:
            return jsonify({"success": False, "error": "prompt 不能为空"}), 400
        
        service = _get_task_service()
        task = service.create_task(
            prompt=prompt,
            cwd=cwd,
            async_mode=async_mode
        )
        
        if async_mode:
            # 异步模式：立即返回任务信息
            return jsonify({
                "success": True,
                "data": task.to_dict(),
                "message": "任务已创建，使用 /claude-tasks/<task_id>/stream 获取实时输出"
            })
        else:
            # 同步模式：等待完成后返回
            return jsonify({
                "success": True,
                "data": {
                    **task.to_dict(),
                    "output": task.output
                }
            })
        
    except Exception as e:
        logger.error(f"创建任务失败: {e}")
        import traceback
        traceback.print_exc()
        return jsonify({"success": False, "error": str(e)}), 500


@copywriting_bp.route('/claude-tasks', methods=['GET'])
def list_claude_tasks():
    """列出所有任务"""
    logger = _get_logger()
    
    try:
        limit = request.args.get('limit', 50, type=int)
        service = _get_task_service()
        tasks = service.list_tasks(limit=limit)
        
        return jsonify({
            "success": True,
            "data": tasks,
            "count": len(tasks)
        })
        
    except Exception as e:
        logger.error(f"获取任务列表失败: {e}")
        return jsonify({"success": False, "error": str(e)}), 500


@copywriting_bp.route('/claude-tasks/<task_id>', methods=['GET'])
def get_claude_task(task_id):
    """获取任务状态"""
    logger = _get_logger()
    
    try:
        service = _get_task_service()
        status = service.get_task_status(task_id)
        
        if not status:
            return jsonify({"success": False, "error": "任务不存在"}), 404
        
        return jsonify({
            "success": True,
            "data": status
        })
        
    except Exception as e:
        logger.error(f"获取任务状态失败: {e}")
        return jsonify({"success": False, "error": str(e)}), 500


@copywriting_bp.route('/claude-tasks/<task_id>/output', methods=['GET'])
def get_claude_task_output(task_id):
    """获取任务完整输出"""
    logger = _get_logger()
    
    try:
        service = _get_task_service()
        output = service.get_task_output(task_id)
        
        if output is None:
            return jsonify({"success": False, "error": "任务不存在"}), 404
        
        # 清理 ANSI 转义序列（可选）
        clean = request.args.get('clean', 'false').lower() == 'true'
        if clean:
            import re
            output = re.sub(r'\x1b\[[0-9;]*[a-zA-Z]', '', output)
            output = re.sub(r'\x1b\][^\x07]*\x07', '', output)
        
        return jsonify({
            "success": True,
            "data": {
                "output": output,
                "length": len(output)
            }
        })
        
    except Exception as e:
        logger.error(f"获取任务输出失败: {e}")
        return jsonify({"success": False, "error": str(e)}), 500


@copywriting_bp.route('/claude-tasks/<task_id>/stream', methods=['GET'])
def stream_claude_task(task_id):
    """
    流式获取任务输出（SSE）
    
    实时推送 Claude 的输出到客户端
    """
    from flask import Response
    logger = _get_logger()
    
    def generate():
        service = _get_task_service()
        timeout = request.args.get('timeout', 300, type=int)
        
        for message in service.stream_output(task_id, timeout=timeout):
            yield f"data: {json.dumps(message, ensure_ascii=False)}\n\n"
    
    return Response(
        generate(),
        mimetype='text/event-stream',
        headers={
            'Cache-Control': 'no-cache',
            'Connection': 'keep-alive',
            'X-Accel-Buffering': 'no'
        }
    )


@copywriting_bp.route('/claude-tasks/<task_id>/cancel', methods=['POST'])
def cancel_claude_task(task_id):
    """取消任务"""
    logger = _get_logger()
    
    try:
        service = _get_task_service()
        result = service.cancel_task(task_id)
        
        return jsonify({
            "success": result,
            "message": "任务已取消" if result else "无法取消任务"
        })
        
    except Exception as e:
        logger.error(f"取消任务失败: {e}")
        return jsonify({"success": False, "error": str(e)}), 500


@copywriting_bp.route('/claude-tasks/cleanup', methods=['POST'])
def cleanup_claude_tasks():
    """清理旧任务"""
    logger = _get_logger()
    
    try:
        data = request.get_json() or {}
        max_age = data.get('max_age_seconds', 3600)
        
        service = _get_task_service()
        count = service.cleanup_old_tasks(max_age_seconds=max_age)
        
        return jsonify({
            "success": True,
            "data": {
                "cleaned": count
            }
        })
        
    except Exception as e:
        logger.error(f"清理任务失败: {e}")
        return jsonify({"success": False, "error": str(e)}), 500


# ============================================================================
# 访客追踪 & 活动日志
# ============================================================================

def _tracking_db():
    _init_services()
    return _services['db']

@copywriting_bp.route('/tracking/visit', methods=['POST'])
def track_visit():
    """记录页面访问"""
    logger = _get_logger()
    try:
        data = request.get_json() or {}
        db = _tracking_db()
        import uuid as _uuid
        visitor_id = str(_uuid.uuid4())

        ip = data.get('ip_address') or request.headers.get('X-Forwarded-For', '').split(',')[0].strip() or request.headers.get('X-Real-IP') or request.remote_addr

        db.log_visitor(
            visitor_id=visitor_id,
            path=data.get('path', '/'),
            user_id=data.get('user_id'),
            ip_address=ip,
            user_agent=data.get('user_agent', ''),
            referer=data.get('referer', ''),
            method=data.get('method', 'GET'),
            status_code=data.get('status_code'),
            device_type=data.get('device_type'),
            browser=data.get('browser'),
            os=data.get('os'),
            session_id=data.get('session_id'),
        )
        return jsonify({'success': True, 'id': visitor_id})
    except Exception as e:
        logger.error(f"记录访问失败: {e}")
        return jsonify({'success': False, 'error': str(e)}), 500


@copywriting_bp.route('/tracking/activity', methods=['POST'])
def track_activity():
    """记录用户活动（单条或批量）"""
    logger = _get_logger()
    try:
        data = request.get_json() or {}
        db = _tracking_db()
        import uuid as _uuid

        if 'events' in data and isinstance(data['events'], list):
            for evt in data['events']:
                if 'id' not in evt:
                    evt['id'] = str(_uuid.uuid4())
            count = db.log_activities_batch(data['events'])
            return jsonify({'success': True, 'recorded': count})

        activity_id = str(_uuid.uuid4())
        db.log_activity(
            activity_id=activity_id,
            action=data.get('action', 'unknown'),
            user_id=data.get('user_id'),
            session_id=data.get('session_id'),
            ip_address=data.get('ip_address'),
            category=data.get('category', 'general'),
            target=data.get('target'),
            target_id=data.get('target_id'),
            details=data.get('details'),
            path=data.get('path'),
        )
        return jsonify({'success': True, 'id': activity_id})
    except Exception as e:
        logger.error(f"记录活动失败: {e}")
        return jsonify({'success': False, 'error': str(e)}), 500


@copywriting_bp.route('/tracking/visitors', methods=['GET'])
def get_visitors():
    """获取访客日志列表"""
    try:
        db = _tracking_db()
        page = request.args.get('page', 1, type=int)
        page_size = request.args.get('page_size', 50, type=int)
        ip = request.args.get('ip')
        path = request.args.get('path')
        user_id = request.args.get('user_id')
        result = db.get_visitor_logs(page=page, page_size=page_size, ip=ip, path=path, user_id=user_id)
        return jsonify({'success': True, **result})
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500


@copywriting_bp.route('/tracking/activities', methods=['GET'])
def get_activities():
    """获取活动日志列表"""
    try:
        db = _tracking_db()
        page = request.args.get('page', 1, type=int)
        page_size = request.args.get('page_size', 50, type=int)
        action = request.args.get('action')
        user_id = request.args.get('user_id')
        result = db.get_activity_logs(page=page, page_size=page_size, action=action, user_id=user_id)
        return jsonify({'success': True, **result})
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500


@copywriting_bp.route('/tracking/stats', methods=['GET'])
def get_tracking_stats():
    """获取访客统计数据"""
    try:
        db = _tracking_db()
        days = request.args.get('days', 30, type=int)
        stats = db.get_visitor_stats(days=days)
        return jsonify({'success': True, 'stats': stats})
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500


@copywriting_bp.route('/tracking/page-trends', methods=['GET'])
def get_page_trends():
    """获取按页面分组的每日访问趋势"""
    try:
        db = _tracking_db()
        days = request.args.get('days', 30, type=int)
        trends = db.get_page_trends(days=days)
        dwell = db.get_page_dwell_stats(days=days)
        return jsonify({'success': True, 'trends': trends, 'dwell_stats': dwell})
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500


@copywriting_bp.route('/tracking/duration', methods=['POST'])
def update_visit_duration():
    """上报页面停留时间"""
    logger = _get_logger()
    try:
        data = request.get_json() or {}
        visit_id = data.get('visit_id')
        duration_ms = data.get('duration_ms')
        if not visit_id or duration_ms is None:
            return jsonify({'success': False, 'error': 'visit_id and duration_ms required'}), 400
        db = _tracking_db()
        ok = db.update_visit_duration(visit_id, int(duration_ms))
        return jsonify({'success': ok})
    except Exception as e:
        logger.error(f"更新停留时间失败: {e}")
        return jsonify({'success': False, 'error': str(e)}), 500


@copywriting_bp.route('/tracking/cleanup', methods=['POST'])
def cleanup_tracking_logs():
    """清理旧日志"""
    try:
        db = _tracking_db()
        data = request.get_json() or {}
        days = data.get('days', 90)
        result = db.cleanup_old_logs(days=days)
        return jsonify({'success': True, **result})
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500


@copywriting_bp.route('/tracking/log-stats', methods=['GET'])
def get_log_stats():
    """获取日志表容量统计（行数、上限、时间范围、数据库大小）"""
    try:
        db = _tracking_db()
        stats = db.get_log_stats()
        return jsonify({'success': True, **stats})
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500


@copywriting_bp.route('/tracking/auto-cleanup', methods=['POST'])
def trigger_auto_cleanup():
    """手动触发自动清理（按保留天数 + 行数上限双重策略）"""
    try:
        db = _tracking_db()
        result = db.auto_cleanup_logs()
        return jsonify({'success': True, **result})
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500
